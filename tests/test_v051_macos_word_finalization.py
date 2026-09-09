from __future__ import annotations

import copy
import contextlib
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
MODULE_PATH = (
    ROOT / "adapters" / "microsoft-word" / "macos" / "word_field_updater.py"
)
SPEC = importlib.util.spec_from_file_location("macos_word_field_updater", MODULE_PATH)
assert SPEC and SPEC.loader
adapter = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(adapter)
REAL_REQUIRE_HOST = adapter._require_host

import finalize_docx  # noqa: E402


def add_field(paragraph, instruction: str, result: str) -> None:
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = result
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


class V051MacosWordFinalizationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.host = patch.object(adapter, "_require_host")
        self.host.start()
        self.addCleanup(self.host.stop)
        self.temp = tempfile.TemporaryDirectory(prefix="v051-macos-word-")
        self.root = Path(self.temp.name).resolve()
        self.input = self.root / "input.docx"
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        add_field(document.add_paragraph(), " QUOTE \"do not update\" ", "unchanged")
        document.sections[0].footer.paragraphs[0].text = ""
        document.save(self.input)
        self.profile = self.root / "profile.json"
        self.profile.write_text("{}", encoding="utf-8")
        self.structure = self.root / "structure.json"
        self.structure.write_text(
            json.dumps({"toc_source": {"approved": False}}), encoding="utf-8"
        )

    def tearDown(self) -> None:
        self.temp.cleanup()

    def request(self, operation: str, **updates):
        request = {
            "protocol_version": "1.1",
            "operation": operation,
            "input_path": str(self.input),
            "profile_path": str(self.profile),
            "structure_map_path": str(self.structure),
            "allowed_field_types": [
                "TOC",
                "PAGE",
                "NUMPAGES",
                "SECTIONPAGES",
                "PAGEREF",
                "REF",
            ],
            "target_software": "microsoft_word",
        }
        if operation == "refresh_fields":
            request["output_path"] = str(self.root / "refreshed.docx")
        if operation == "verify_only":
            request["pdf_output_path"] = str(self.root / "verification.pdf")
        request.update(updates)
        return request

    @staticmethod
    def snapshot(page_count: int = 1, result: str = "1"):
        return [
            [],
            [],
            [[0, 1, page_count, page_count, 1, True, 1, "decimal"]],
            page_count,
            [["PAGE", " PAGE ", result]],
            [],
        ]

    def response(self, operation: str, snapshots=None, **updates):
        if snapshots is None:
            snapshot = self.snapshot()
            snapshots = (
                [snapshot, snapshot]
                if operation in {"refresh_fields", "verify_only"}
                else [snapshot]
            )
        raw = [
            "success",
            operation,
            "16.synthetic",
            operation == "refresh_fields",
            operation != "refresh_fields",
            operation == "verify_only",
            "exact_document_closed_without_save",
            True,
            0,
            snapshots[-1][3],
            len(snapshots[-1][0]),
            1,
            [["PAGE", 1]] if operation == "refresh_fields" else [],
            snapshots,
            [True, False, False, False]
            if operation == "verify_only"
            else [True],
            [],
        ]
        indexes = {
            "docx_saved": 3,
            "read_only": 4,
            "pdf_exported": 5,
            "close_outcome": 6,
            "controls_restored": 7,
            "structural_changes": 8,
            "page_count": 9,
            "toc_count": 10,
            "verified_count": 11,
            "updated_pairs": 12,
            "saved_observations": 14,
            "story_observations": 15,
        }
        for name, value in updates.items():
            raw[indexes[name]] = value
        return raw

    def completed(self, operation: str, *, before_return=None, response=None):
        raw = self.response(operation) if response is None else response

        def execute(command, **kwargs):
            self.assertEqual(adapter.OSASCRIPT, command[0])
            self.assertEqual(str(adapter.APPLESCRIPT), command[1])
            self.assertEqual(operation, command[2])
            self.assertNotIn("QUOTE", command[5])
            if before_return:
                before_return(command)
            return subprocess.CompletedProcess(command, 0, json.dumps(raw), "")

        return execute

    @staticmethod
    def story_observations(plan):
        return [
            [
                item["section_index"],
                item["story_label"],
                item["ownership"],
                len(item["field_identities"]),
                item["field_identities"],
            ]
            for item in plan
        ]

    def rewrite_input(self, transform, *, drop=()) -> None:
        replacement = self.root / "rewritten.docx"
        with zipfile.ZipFile(self.input) as source, zipfile.ZipFile(
            replacement, "w", zipfile.ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                if info.filename in drop:
                    continue
                target.writestr(
                    info,
                    transform(info.filename, source.read(info.filename)),
                )
        replacement.replace(self.input)

    def test_primary_footer_preflight_rejects_missing_before_word(self) -> None:
        private_text = "private-synthetic-footer-preflight"
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        document.add_paragraph(private_text)
        document.save(self.input)
        with patch.object(adapter.subprocess, "run") as run:
            with self.assertRaises(adapter.AdapterError) as caught:
                adapter.run_request(self.request("measure_layout"))
        run.assert_not_called()
        self.assertIn("primary footer", str(caught.exception))
        self.assertNotIn(str(self.input), str(caught.exception))
        self.assertNotIn(private_text, str(caught.exception))

    def test_primary_footer_preflight_rejects_mid_chain_broken_owner(self) -> None:
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        document.sections[0].footer.paragraphs[0].text = ""
        second = document.add_section(WD_SECTION.NEW_PAGE)
        second.footer.is_linked_to_previous = False
        second.footer.paragraphs[0].text = ""
        document.save(self.input)

        def break_second_owner(name, data):
            if name != "word/document.xml":
                return data
            root = etree.fromstring(data)
            sections = root.xpath(".//w:sectPr", namespaces=adapter.NS)
            reference = sections[1].xpath(
                "./w:footerReference[@w:type='default']",
                namespaces=adapter.NS,
            )[0]
            del reference.attrib[qn("r:id")]
            return etree.tostring(
                root, xml_declaration=True, encoding="UTF-8", standalone=True
            )

        self.rewrite_input(break_second_owner)
        with patch.object(adapter.subprocess, "run") as run:
            with self.assertRaisesRegex(adapter.AdapterError, "primary footer"):
                adapter.run_request(self.request("measure_layout"))
        run.assert_not_called()

    def test_primary_footer_preflight_rejects_wrong_relationship_and_missing_part(
        self,
    ) -> None:
        with zipfile.ZipFile(self.input) as package:
            document = etree.fromstring(package.read("word/document.xml"))
            rel_id = document.xpath(
                ".//w:sectPr/w:footerReference[@w:type='default']/@r:id",
                namespaces=adapter.NS,
            )[0]
            relationships = etree.fromstring(
                package.read("word/_rels/document.xml.rels")
            )
            footer_part = adapter.posixpath.normpath(
                adapter.posixpath.join(
                    "word",
                    next(item for item in relationships if item.get("Id") == rel_id).get(
                        "Target"
                    ),
                )
            )

        for label in ("wrong_relationship", "missing_part"):
            original = self.input.read_bytes()

            def mutate(name, data, *, current=label):
                if (
                    current == "wrong_relationship"
                    and name == "word/_rels/document.xml.rels"
                ):
                    root = etree.fromstring(data)
                    relationship = next(
                        item for item in root if item.get("Id") == rel_id
                    )
                    relationship.set(
                        "Type",
                        adapter.STORY_RELATIONSHIP_TYPES["header"],
                    )
                    return etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                return data

            self.rewrite_input(
                mutate,
                drop=(footer_part,) if label == "missing_part" else (),
            )
            with self.subTest(label=label), patch.object(
                adapter.subprocess, "run"
            ) as run:
                with self.assertRaises(adapter.AdapterError) as caught:
                    adapter.run_request(self.request("measure_layout"))
                run.assert_not_called()
                self.assertIn("primary footer", str(caught.exception))
                self.assertNotIn(str(self.input), str(caught.exception))
            self.input.write_bytes(original)

    def test_shared_inherited_primary_footer_passes_preflight(self) -> None:
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        document.sections[0].footer.paragraphs[0].text = ""
        second = document.add_section(WD_SECTION.NEW_PAGE)
        self.assertTrue(second.footer.is_linked_to_previous)
        document.save(self.input)
        snapshot = self.snapshot(page_count=2)
        snapshot[2] = [
            [0, 1, 1, 1, 1, True, 1, "decimal"],
            [1, 2, 2, 2, 2, False, 0, "decimal"],
        ]
        response = self.response("measure_layout", snapshots=[snapshot])
        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("measure_layout", response=response),
        ) as run:
            result = adapter.run_request(self.request("measure_layout"))
        run.assert_called_once()
        self.assertEqual(2, result["page_count"])

    def test_story_access_plan_uses_only_explicit_field_bearing_owners(self) -> None:
        self.assertEqual([], adapter._story_access_plan(self.input))
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        first = document.sections[0]
        localized = document.styles.add_style(
            "本地化合成页眉", WD_STYLE_TYPE.PARAGRAPH
        )
        first.header.paragraphs[0].style = localized
        add_field(first.header.paragraphs[0], " NUMPAGES ", "2")
        first.different_first_page_header_footer = True
        add_field(first.first_page_header.paragraphs[0], " REF Synthetic ", "1")
        document.settings.odd_and_even_pages_header_footer = True
        add_field(first.even_page_footer.paragraphs[0], " PAGE ", "1")
        second = document.add_section(WD_SECTION.NEW_PAGE)
        self.assertTrue(second.header.is_linked_to_previous)
        second.footer.is_linked_to_previous = False
        add_field(second.footer.paragraphs[0], " SECTIONPAGES ", "1")
        document.save(self.input)

        plan = adapter._story_access_plan(self.input)
        self.assertEqual(4, len(plan))
        owners = {
            (item["section_index"], item["story_label"]): item
            for item in plan
        }
        self.assertEqual(
            {
                (1, "header_primary"),
                (1, "header_first"),
                (1, "footer_even"),
                (2, "footer_primary"),
            },
            set(owners),
        )
        self.assertEqual("first_section", owners[(1, "header_primary")]["ownership"])
        self.assertEqual("independent", owners[(2, "footer_primary")]["ownership"])
        self.assertEqual(
            [["NUMPAGES", "NUMPAGES"]],
            owners[(1, "header_primary")]["field_identities"],
        )
        self.assertNotIn((2, "header_primary"), owners)
        token = adapter._encode_story_access_plan(plan)
        self.assertNotIn("NUMPAGES", token)
        self.assertNotIn("REF", token)

    def test_field_bearing_story_without_valid_owner_is_rejected_before_word(self) -> None:
        replacement = self.root / "unowned.docx"
        with zipfile.ZipFile(self.input) as source, zipfile.ZipFile(
            replacement, "w", zipfile.ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                target.writestr(info, source.read(info.filename))
            header = OxmlElement("w:hdr")
            paragraph = OxmlElement("w:p")
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), " PAGE ")
            paragraph.append(field)
            header.append(paragraph)
            target.writestr(
                "word/header99.xml",
                etree.tostring(
                    header,
                    xml_declaration=True,
                    encoding="UTF-8",
                    standalone=True,
                ),
            )
        replacement.replace(self.input)
        with patch.object(adapter.subprocess, "run") as run:
            with self.assertRaisesRegex(adapter.AdapterError, "no explicit owner"):
                adapter.run_request(self.request("measure_layout"))
        run.assert_not_called()

    def test_story_access_observation_mismatch_fails_closed(self) -> None:
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        secret_instruction = " REF ConfidentialSyntheticBookmark "
        add_field(document.sections[0].footer.paragraphs[0], secret_instruction, "1")
        document.save(self.input)
        plan = adapter._story_access_plan(self.input)
        expected = self.story_observations(plan)
        response = self.response(
            "measure_layout", verified_count=2, story_observations=expected
        )
        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("measure_layout", response=response),
        ) as run:
            result = adapter.run_request(self.request("measure_layout"))
        self.assertEqual(adapter._encode_story_access_plan(plan), run.call_args.args[0][6])
        self.assertNotIn("story_access_plan", result)
        self.assertNotIn("story_observations", result)
        self.assertNotIn("ConfidentialSyntheticBookmark", json.dumps(result))
        self.assertNotIn("ConfidentialSyntheticBookmark", run.call_args.args[0][6])

        cases = {}
        wrong = json.loads(json.dumps(expected))
        wrong[0][0] += 1
        cases["owner"] = wrong
        wrong = json.loads(json.dumps(expected))
        wrong[0][2] = "independent"
        cases["ownership"] = wrong
        wrong = json.loads(json.dumps(expected))
        wrong[0][3] += 1
        cases["count"] = wrong
        wrong = json.loads(json.dumps(expected))
        wrong[0][4][0][1] = " NUMPAGES "
        cases["instruction"] = wrong
        cases["extra"] = expected + [expected[0]]
        for label, wrong in cases.items():
            response = self.response(
                "measure_layout", verified_count=2, story_observations=wrong
            )
            with self.subTest(label=label), patch.object(
                adapter.subprocess,
                "run",
                side_effect=self.completed("measure_layout", response=response),
            ):
                with self.assertRaisesRegex(adapter.AdapterError, "story access|story field"):
                    adapter.run_request(self.request("measure_layout"))

    def test_invalid_story_relationship_is_rejected_before_word(self) -> None:
        document = Document()
        add_field(document.add_paragraph(), " PAGE ", "1")
        add_field(document.sections[0].footer.paragraphs[0], " PAGE ", "1")
        document.save(self.input)
        replacement = self.root / "invalid-relationship.docx"
        with zipfile.ZipFile(self.input) as source, zipfile.ZipFile(
            replacement, "w", zipfile.ZIP_DEFLATED
        ) as target:
            for info in source.infolist():
                data = source.read(info.filename)
                if info.filename == "word/_rels/document.xml.rels":
                    root = etree.fromstring(data)
                    relationship = next(
                        item
                        for item in root
                        if str(item.get("Type", "")).endswith("/footer")
                    )
                    relationship.set(
                        "Type",
                        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/header",
                    )
                    data = etree.tostring(
                        root,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    )
                target.writestr(info, data)
        replacement.replace(self.input)
        with patch.object(adapter.subprocess, "run") as run:
            with self.assertRaisesRegex(adapter.AdapterError, "primary footer"):
                adapter.run_request(self.request("measure_layout"))
        run.assert_not_called()

    def test_three_protocol_operations_are_exact_path_owned_and_bounded(self) -> None:
        input_hash = adapter.sha256(self.input)
        with patch.object(adapter.subprocess, "run", side_effect=self.completed("measure_layout")) as run:
            result = adapter.run_request(self.request("measure_layout"))
        self.assertEqual("measure_layout", result["operation"])
        self.assertTrue(result["read_only_verified"])
        self.assertFalse(result["saved"])
        self.assertEqual(str(self.input), run.call_args.args[0][3])
        self.assertEqual(adapter.TIMEOUT_SECONDS, run.call_args.kwargs["timeout"])

        with patch.object(adapter.subprocess, "run", side_effect=self.completed("refresh_fields")) as run:
            result = adapter.run_request(self.request("refresh_fields"))
        self.assertTrue(result["saved"])
        self.assertTrue(result["field_cache_verified"])
        self.assertEqual(str(self.root / "refreshed.docx"), run.call_args.args[0][3])
        self.assertEqual(input_hash, adapter.sha256(self.input))

        def export_pdf(command):
            document = pymupdf.open()
            document.new_page()
            document.save(command[4])
            document.close()

        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("verify_only", before_return=export_pdf),
        ):
            result = adapter.run_request(self.request("verify_only"))
        self.assertTrue(result["pdf_exported"])
        self.assertFalse(result["saved"])
        self.assertEqual([True, False, False, False], result["saved_observations"])
        self.assertEqual(input_hash, adapter.sha256(self.input))

    def test_existing_core_protocol_accepts_all_adapter_operation_shapes(self) -> None:
        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("measure_layout"),
        ):
            measured = adapter.run_request(self.request("measure_layout"))
        with patch.object(
            finalize_docx,
            "_invoke_external_command",
            return_value=subprocess.CompletedProcess(
                [], 0, json.dumps(measured), ""
            ),
        ):
            accepted = finalize_docx.external_measure(
                self.input,
                json.dumps(["synthetic-adapter"]),
                self.profile,
                self.structure,
                "microsoft_word",
            )
        self.assertEqual("external", accepted["backend"])
        self.assertEqual(adapter.BACKEND, accepted["implementation_backend"])
        self.assertEqual("external", finalize_docx.canonical_backend_projection(accepted)["backend"])

        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("refresh_fields"),
        ):
            refreshed = adapter.run_request(self.request("refresh_fields"))
        refreshed_path = self.root / "refreshed.docx"
        with patch.object(
            finalize_docx,
            "_invoke_external_command",
            return_value=subprocess.CompletedProcess(
                [], 0, json.dumps(refreshed), ""
            ),
        ):
            accepted = finalize_docx.external_refresh(
                self.input,
                refreshed_path,
                json.dumps(["synthetic-adapter"]),
                self.profile,
                self.structure,
                None,
                "microsoft_word",
            )
        self.assertTrue(accepted["field_cache_verified"])
        self.assertEqual("external", accepted["backend"])
        self.assertEqual(adapter.BACKEND, accepted["implementation_backend"])
        self.assertEqual("external", finalize_docx.canonical_backend_projection(accepted)["backend"])

        def export_pdf(command):
            document = pymupdf.open()
            document.new_page()
            document.save(command[4])
            document.close()

        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("verify_only", before_return=export_pdf),
        ):
            verified = adapter.run_request(self.request("verify_only"))
        pdf = self.root / "verification.pdf"
        with patch.object(
            finalize_docx,
            "_invoke_external_command",
            return_value=subprocess.CompletedProcess(
                [], 0, json.dumps(verified), ""
            ),
        ):
            accepted = finalize_docx.external_verify(
                self.input,
                json.dumps(["synthetic-adapter"]),
                self.profile,
                self.structure,
                pdf,
                "microsoft_word",
                expected_page_count=1,
            )
        self.assertTrue(accepted["read_only_verified"])
        self.assertEqual("external", accepted["backend"])
        self.assertEqual(adapter.BACKEND, accepted["implementation_backend"])
        self.assertEqual("external", finalize_docx.canonical_backend_projection(accepted)["backend"])

        operations = [
            (measured, lambda: finalize_docx.external_measure(self.input, "synthetic", self.profile, self.structure, "microsoft_word")),
            (refreshed, lambda: finalize_docx.external_refresh(self.input, refreshed_path, "synthetic", self.profile, self.structure, None, "microsoft_word")),
            (verified, lambda: finalize_docx.external_verify(self.input, "synthetic", self.profile, self.structure, pdf, "microsoft_word", expected_page_count=1)),
        ]
        for response, call in operations:
            for label in (None, "external", "unknown_adapter", "libreoffice_uno", 7, {}):
                value = copy.deepcopy(response)
                if label is None:
                    value.pop("backend")  # Existing omitted-label protocol remains valid.
                else:
                    value["backend"] = label
                with self.subTest(operation=value["operation"], label=label), patch.object(
                    finalize_docx, "_invoke_external_command",
                    return_value=subprocess.CompletedProcess([], 0, json.dumps(value), ""),
                ):
                    if label is None or label == "external":
                        result = call()
                        self.assertEqual("external", result["backend"])
                        self.assertEqual("external", result["implementation_backend"])
                    else:
                        result = call()
                        self.assertEqual(label, result["backend"])
                        self.assertNotIn("implementation_backend", result)
                        if label != "libreoffice_uno":
                            with self.assertRaises((ValueError, TypeError)):
                                finalize_docx.canonical_backend_projection(result)
                        from field_completion import final_ready_evidence_errors
                        errors = final_ready_evidence_errors({"delivery_status": "selective_verified", "backend": result["backend"]})
                        self.assertTrue(any(error.startswith("backend=") for error in errors))
            for mutation in ({"software": "LibreOffice"}, {"status": "error"}, {"saved": not response["saved"]}):
                value = dict(response, **mutation)
                with self.subTest(operation=value["operation"], mutation=mutation), patch.object(
                    finalize_docx, "_invoke_external_command",
                    return_value=subprocess.CompletedProcess([], 0, json.dumps(value), ""),
                ):
                    with self.assertRaises(finalize_docx.FormatMonographError):
                        call()

    def test_real_adapter_protocol_reaches_finalization_publication(self) -> None:
        self._assert_adapter_protocol_publication("stale")

    def test_first_toc_adapter_protocol_reaches_finalization_publication(self) -> None:
        self._assert_adapter_protocol_publication("code_only")

    def test_no_toc_cached_scalars_reach_publication_only_with_complete_word_evidence(self) -> None:
        self._assert_adapter_protocol_publication("refreshed", no_toc=True)

    def test_bad_duplicate_footer_cannot_publish_or_reach_verify_status_completion(self) -> None:
        self._assert_adapter_protocol_publication("refreshed", no_toc=True, duplicate_footer=True)

    def _assert_adapter_protocol_publication(self, input_cache_status: str, *, no_toc: bool = False, duplicate_footer: bool = False) -> None:
        from test_v11_execution import approved_v11_profile
        from test_v024_finalization import field_char, run_element
        from field_writeback import parse_fields, _result_text_nodes
        from structure_map import candidate_structure_map
        from field_completion import (
            completion_evidence, final_ready_evidence_errors,
            finalization_evidence_shape_errors,
        )

        document = Document()
        document.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
        toc = document.add_paragraph(style="TOC 1" if input_cache_status == "stale" else "Normal")
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), "SyntheticHeading")
        link.append(run_element("w:t", "Synthetic heading"))
        toc._p.extend([
            field_char("begin"), run_element("w:instrText", ' TOC \\o "1-1" \\h \\z '),
            field_char("separate"), link,
            run_element("w:tab", ""), run_element("w:t", "0"), field_char("end"),
        ])
        toc._p.xpath(".//w:fldChar")[0].set(qn("w:dirty"), "true")
        heading = document.add_paragraph("Synthetic heading", style="Heading 1")
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), "1")
        start.set(qn("w:name"), "SyntheticHeading")
        heading._p.insert(0, start)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), "1")
        heading._p.append(end)
        add_field(document.add_paragraph(), " PAGE ", "0")
        document.sections[0].footer.paragraphs[0].text = ""
        if no_toc:
            toc._p.getparent().remove(toc._p)
            add_field(document.add_paragraph(), " REF SyntheticHeading ", "0")
            add_field(document.sections[0].footer.paragraphs[0], " SECTIONPAGES ", "1")
        if duplicate_footer:
            for i in range(2):
                add_field(document.sections[0].footer.add_paragraph(f"Distinct footer {i} "), " PAGE ", "9")
        document.save(self.input)
        profile = approved_v11_profile()
        profile["rules"] = [r for r in profile["rules"] if r["selector"]["kind"] == "paragraph_role"]
        self.profile.write_text(json.dumps(profile))
        structure = candidate_structure_map(self.input)
        structure["status"] = "approved"
        for item in structure["headings"]:
            item["approved"] = True
        for item in structure["paragraph_roles"]:
            item["approved"] = item["role"] != "unknown"
        structure["toc_source"].update(approved=not no_toc, mode="heading_styles", levels=1)
        self.structure.write_text(json.dumps(structure))
        applied = subprocess.run(
            [sys.executable, str(ROOT / "format-monograph/scripts/apply_profile.py"),
             str(self.input), "--profile", str(self.profile), "--structure-map", str(self.structure),
             "--output-dir", str(self.root / "applied"), "--allow-missing-fonts"],
            capture_output=True, text=True,
        )
        self.assertEqual(0, applied.returncode, applied.stderr)
        formatted = self.root / "applied/input-formatted.docx"
        self.assertEqual(input_cache_status, finalize_docx.field_cache_inventory(formatted)["status"])
        allowed = sorted(adapter.ALLOWED_FIELD_TYPES - {"="})
        self.assertEqual(not no_toc, "TOC" in adapter._effective_allowed(allowed, formatted, structure))
        unapproved = copy.deepcopy(structure)
        unapproved["toc_source"]["approved"] = False
        self.assertNotIn("TOC", adapter._effective_allowed(allowed, formatted, unapproved))

        cases = ("external", "auto", "auto_verify_failure", "auto_copy") if no_toc else ("external", "auto", "auto_verify_failure")
        if duplicate_footer:
            cases = ("external",)
        for case in cases:
            mode = "external" if case == "external" else "auto"
            with self.subTest(case=case):
                folder = self.root / case
                folder.mkdir()
                operations = []
                raw_responses = []

                def word(command, **kwargs):
                    operation = command[2]
                    path = Path(command[3])
                    if operation == "refresh_fields":
                        doc = Document(path)
                        for record in parse_fields(doc._element):
                            marker = record.simple if record.form == "simple" else record.begin
                            marker.attrib.pop(qn("w:dirty"), None)
                            _result_text_nodes(doc._element, record)[-1].text = "1"
                        if input_cache_status == "code_only":
                            doc.paragraphs[0].style = doc.styles["TOC 1"]
                        doc.save(path)
                    if operation == "verify_only":
                        with pymupdf.open() as pdf:
                            pdf.new_page()
                            pdf.save(command[4])
                    snapshot = self.snapshot()
                    snapshot[0] = [] if no_toc else ["Synthetic heading\t1\r"]
                    snapshot[1] = [] if no_toc else [1]
                    snapshots = [snapshot] if operation == "measure_layout" else [snapshot, snapshot]
                    if duplicate_footer:
                        snapshot[4] = [["PAGE", " PAGE ", "1"], ["REF", " REF SyntheticHeading ", "1"],
                                       ["SECTIONPAGES", " SECTIONPAGES ", "1"], ["PAGE", " PAGE ", "1"], ["PAGE", " PAGE ", "1"]]
                        snapshot.append([self.story_observations(adapter._story_access_plan(path)),
                                         [[1, "footer_primary", "PAGE", " PAGE ", "1"]] * 2])
                    result = self.response(
                        operation, snapshots, verified_count=5 if duplicate_footer else (3 if no_toc else 2),
                        updated_pairs=([["PAGE", 3 if duplicate_footer else 1], ["REF", 1], ["SECTIONPAGES", 1]] if no_toc else [["PAGE", 1], ["TOC", 1]]) if operation == "refresh_fields" else [],
                        story_observations=self.story_observations(adapter._story_access_plan(path)),
                    )
                    return subprocess.CompletedProcess(command, 0, json.dumps(result), "")

                def invoke(command, request, label):
                    operations.append(request["operation"])
                    with patch.object(adapter.subprocess, "run", side_effect=word):
                        try:
                            response = adapter.run_request(request)
                        except adapter.AdapterError as exc:
                            return subprocess.CompletedProcess([], 1, "", str(exc))
                    raw_responses.append(response)
                    if case == "auto_verify_failure" and request["operation"] == "verify_only":
                        response["read_only_verified"] = False
                    return subprocess.CompletedProcess([], 0, json.dumps(response), "")

                argv = ["finalize_docx.py", str(formatted), "--source", str(self.input),
                        "--profile", str(self.profile), "--structure-map", str(self.structure),
                        "--output", str(folder / "final.docx"), "--status-output", str(folder / "finalization.json"),
                        "--pdf-output", str(folder / "target.pdf"), "--field-updater", mode,
                        "--target-software", "microsoft_word"]
                if case != "auto_copy":
                    argv.extend(["--field-updater-command", "synthetic-adapter"])
                if case == "auto_verify_failure":
                    argv.append("--approve-deferred")
                stderr = io.StringIO()
                if duplicate_footer:
                    import run_monograph as runtime
                    from types import SimpleNamespace
                    state = {
                        "status": "candidate_ready",
                        "source": {"path": str(self.input), "sha256": adapter.sha256(self.input)},
                        "profile": {"path": str(self.profile), "sha256": adapter.sha256(self.profile)},
                        "structure_map": {"path": str(self.structure), "sha256": adapter.sha256(self.structure)},
                        "artifacts": {"formatted": str(formatted)}, "blockers": [],
                        "qa_groups": [], "frozen_scopes": [], "stages": {}, "metrics": {},
                    }
                    args = SimpleNamespace(work_dir=folder, resume=False, field_updater="external",
                                           field_updater_command="synthetic-adapter", target_software="microsoft_word",
                                           renderer=None, approve_deferred=False, json=True)
                    def run(script, *arguments):
                        stdout, error = io.StringIO(), io.StringIO()
                        with patch.object(sys, "argv", [script, *map(str, arguments)]), patch.object(
                            finalize_docx, "_invoke_external_command", side_effect=invoke
                        ), contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(error):
                            code = finalize_docx.main()
                        self.assertIn("footer cache", error.getvalue())
                        return subprocess.CompletedProcess([], code, stdout.getvalue(), error.getvalue())
                    with patch.object(runtime, "load_state", return_value=state), patch.object(runtime, "save_state"), patch.object(
                        runtime, "run_script", side_effect=run
                    ) as scripts, contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(stderr):
                        self.assertEqual(2, runtime.finalize(args))
                        self.assertEqual(["measure_layout", "refresh_fields"], operations)
                        self.assertEqual(1, scripts.call_count)
                        self.assertFalse((folder / "final/finalization.json").exists())
                        self.assertNotIn("finalized", state["artifacts"])
                        self.assertEqual(2, runtime.verify(args))
                        self.assertEqual(0, runtime.status(args))
                        self.assertEqual(1, scripts.call_count)  # verify/status cannot execute or upgrade it.
                    self.assertEqual("candidate_ready", state["status"])
                    self.assertTrue(state["field_writeback"]["completion_evidence_errors"])
                    continue
                with patch.object(sys, "argv", argv), patch.object(
                    finalize_docx, "_invoke_external_command", side_effect=invoke
                ), contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(stderr):
                    self.assertEqual(1 if case == "auto_verify_failure" else 0, finalize_docx.main(), stderr.getvalue())
                self.assertEqual([] if case == "auto_copy" else ["measure_layout", "refresh_fields", "verify_only"], operations)
                self.assertTrue(all(r["backend"] == adapter.BACKEND for r in raw_responses))
                evidence_path = folder / "finalization.json"
                if case == "auto_verify_failure":
                    # Existing failed-verification PDF is not valid deferred
                    # publication evidence. Preserve the rejection, not a GO.
                    self.assertFalse(evidence_path.exists())
                    self.assertIn("status target PDF binding is missing", stderr.getvalue())
                    evidence_path = next(folder.glob(".format-monograph-finalize-*/finalization.json"))
                evidence = json.loads(evidence_path.read_text())
                self.assertEqual([], finalization_evidence_shape_errors(evidence))
                if case == "auto_copy":
                    self.assertEqual("refreshed", evidence["delivery_field_status"])
                    self.assertEqual("not_needed", evidence["field_backend"]["backend"])
                    self.assertFalse(evidence["field_completion"]["final_ready_eligible"])
                    self.assertNotEqual("no_fields", evidence["field_completion"]["completion_scope"])
                    self.assertTrue(final_ready_evidence_errors(completion_evidence(evidence)))
                    continue
                if case == "auto_verify_failure":
                    self.assertEqual("deferred", evidence["delivery_field_status"])
                    self.assertFalse(evidence["field_completion"]["final_ready_eligible"])
                    self.assertTrue(final_ready_evidence_errors(completion_evidence(evidence)))
                    self.assertEqual("external", evidence["field_backend"]["attempt"]["backend"])
                    audit = json.loads((evidence_path.parent / "finalization-backend-audit.json").read_text())["backend"]
                    self.assertEqual(adapter.BACKEND, audit["attempted_backend"]["implementation_backend"])
                    self.assertEqual("external_field_workflow", audit["attempted_backend"]["failure"]["stage"])
                    continue
                self.assertEqual([], final_ready_evidence_errors(completion_evidence(evidence)))
                self.assertEqual(input_cache_status, evidence["input_field_cache"]["status"])
                self.assertEqual("refreshed", evidence["output_field_cache"]["status"])
                self.assertTrue(evidence["field_completion"]["field_gate_completed"])
                self.assertTrue(evidence["field_completion"]["final_ready_eligible"])
                self.assertEqual({"status": "pass", "errors": []}, evidence["field_completion"]["evidence_validation"])
                self.assertEqual("external", evidence["field_backend"]["backend"])
                self.assertEqual("pass", evidence["content_integrity"])
                self.assertEqual("pass", evidence["protected_object_integrity"])
                self.assertEqual("pass", evidence["effective_font_integrity"])
                audit = json.loads((folder / "finalization-backend-audit.json").read_text())["backend"]
                for response in [audit, *audit["layout_measurements"], audit["read_only_verification"]]:
                    self.assertEqual("external", response["backend"])
                    self.assertEqual(adapter.BACKEND, response["implementation_backend"])
                for key in ("finalized_docx", "word_verification_pdf"):
                    binding = evidence["artifact_binding"][key]
                    self.assertEqual(adapter.sha256(Path(binding["path"])), binding["sha256"])

    def test_non_macos_or_missing_osascript_host_is_unavailable(self) -> None:
        with patch.object(adapter.sys, "platform", "linux"):
            with self.assertRaisesRegex(adapter.AdapterError, "requires executable"):
                REAL_REQUIRE_HOST()

    def test_whitelist_and_toc_require_structure_approval(self) -> None:
        with patch.object(adapter.subprocess, "run", side_effect=self.completed("measure_layout")) as run:
            adapter.run_request(self.request("measure_layout"))
        token = run.call_args.args[0][5]
        self.assertNotIn("TOC", token)
        self.assertNotIn("QUOTE", token)
        self.assertEqual(
            {"PAGE", "NUMPAGES", "SECTIONPAGES", "PAGEREF", "REF"},
            set(token.split("|")),
        )

        self.structure.write_text(
            json.dumps({"toc_source": {"approved": True}}), encoding="utf-8"
        )
        with patch.object(
            adapter, "audit_structure_toc_source_operations", return_value=[]
        ), patch.object(adapter, "load_document", return_value=object()), patch.object(
            adapter.subprocess, "run", side_effect=self.completed("measure_layout")
        ) as run:
            adapter.run_request(self.request("measure_layout"))
        self.assertIn("TOC", run.call_args.args[0][5].split("|"))

        with patch.object(
            adapter,
            "audit_structure_toc_source_operations",
            return_value=[{"reason": "toc_source_contract_mismatch"}],
        ), patch.object(adapter, "load_document", return_value=object()), patch.object(
            adapter.subprocess, "run"
        ) as run:
            with self.assertRaisesRegex(adapter.AdapterError, "approved structure"):
                adapter.run_request(self.request("measure_layout"))
        run.assert_not_called()

    def test_two_equal_rounds_or_equal_second_and_third_are_required(self) -> None:
        first = self.snapshot(1, "1")
        second = self.snapshot(2, "2")
        third = self.snapshot(3, "3")
        accepted = self.response("refresh_fields", [first, second, second])
        accepted[9] = 2
        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("refresh_fields", response=accepted),
        ):
            result = adapter.run_request(self.request("refresh_fields"))
        self.assertEqual(3, len(result["convergence"]))

        for snapshots in ([first], [first, second], [first, second, third], [first, first, first, first]):
            output = self.root / "refreshed.docx"
            output.unlink(missing_ok=True)
            raw = self.response("refresh_fields", snapshots)
            raw[9] = snapshots[-1][3]
            with self.subTest(rounds=len(snapshots)), patch.object(
                adapter.subprocess,
                "run",
                side_effect=self.completed("refresh_fields", response=raw),
            ):
                with self.assertRaisesRegex(adapter.AdapterError, "converge"):
                    adapter.run_request(self.request("refresh_fields"))
            self.assertFalse(output.exists())

    def test_read_only_pdf_missing_invalid_mismatch_and_save_contradiction_fail(self) -> None:
        cases = {
            "missing": (None, self.response("verify_only"), "did not create"),
            "invalid": (lambda command: Path(command[4]).write_bytes(b"not pdf"), self.response("verify_only"), "unreadable"),
            "saved": (None, self.response("verify_only", docx_saved=True), "contradictory"),
        }
        for label, (producer, raw, message) in cases.items():
            pdf = self.root / "verification.pdf"
            pdf.unlink(missing_ok=True)
            with self.subTest(label=label), patch.object(
                adapter.subprocess,
                "run",
                side_effect=self.completed("verify_only", before_return=producer, response=raw),
            ):
                with self.assertRaisesRegex(adapter.AdapterError, message):
                    adapter.run_request(self.request("verify_only"))
            self.assertFalse(pdf.exists())

        def two_page_pdf(command):
            document = pymupdf.open()
            document.new_page()
            document.new_page()
            document.save(command[4])
            document.close()

        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("verify_only", before_return=two_page_pdf),
        ):
            with self.assertRaisesRegex(adapter.AdapterError, "page count"):
                adapter.run_request(self.request("verify_only"))

    def test_dialog_timeout_restore_close_and_structural_evidence_fail_closed(self) -> None:
        failures = (
            ({"controls_restored": False}, "bounded operation"),
            ({"close_outcome": "close_not_verified"}, "bounded operation"),
            ({"structural_changes": 1}, "bounded operation"),
        )
        for updates, message in failures:
            raw = self.response("measure_layout", **updates)
            with self.subTest(updates=updates), patch.object(
                adapter.subprocess,
                "run",
                side_effect=self.completed("measure_layout", response=raw),
            ):
                with self.assertRaisesRegex(adapter.AdapterError, message):
                    adapter.run_request(self.request("measure_layout"))

        failed = subprocess.CompletedProcess([], 1, "", "Word dialog blocked automation")
        with patch.object(adapter.subprocess, "run", return_value=failed):
            with self.assertRaisesRegex(adapter.AdapterError, "dialog blocked"):
                adapter.run_request(self.request("measure_layout"))
        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=subprocess.TimeoutExpired(
                [adapter.OSASCRIPT], adapter.TIMEOUT_SECONDS
            ),
        ):
            with self.assertRaisesRegex(adapter.AdapterError, "restoration are unconfirmed"):
                adapter.run_request(self.request("measure_layout"))

    def test_timeout_budgets_leave_ordered_cleanup_headroom(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        run_body = source.split("on run argv", 1)[1]
        main_timeout = run_body.index("with timeout of 480 seconds")
        main_timeout_end = run_body.index("end timeout", main_timeout)
        error_handler = run_body.index("on error failureMessage", main_timeout_end)
        cleanup_close = run_body.index(
            "set closeOutcome to my closeExactDocument(inputPath)", error_handler
        )
        cleanup_restore = run_body.index(
            "set restoreErrors to my restoreControls", cleanup_close
        )
        self.assertLess(main_timeout_end, error_handler)
        self.assertLess(error_handler, cleanup_close)
        self.assertLess(cleanup_close, cleanup_restore)
        self.assertEqual(570, adapter.TIMEOUT_SECONDS)
        self.assertGreater(adapter.TIMEOUT_SECONDS, 480 + 30 + 30)
        self.assertLess(adapter.TIMEOUT_SECONDS, 600)

        adapter_source = MODULE_PATH.read_text(encoding="utf-8")
        finalize_source = Path(finalize_docx.__file__).read_text(encoding="utf-8")
        self.assertIn("timeout=TIMEOUT_SECONDS", adapter_source)
        self.assertIn("timeout=600", finalize_source)

    def test_each_pdf_substage_failure_retains_location_and_discards_pdf(self) -> None:
        for stage in (
            "verify.pdf_export",
            "verify.pdf_export.reacquire",
            "verify.pdf_export.read_only",
            "verify.pdf_export.saved",
        ):
            with self.subTest(stage=stage):
                original = self.input.read_bytes()
                pdf = self.root / "verification.pdf"

                def fail(command, **_kwargs):
                    # A PDF may already exist even when the whole operation fails.
                    Path(command[4]).write_bytes(b"synthetic partial PDF")
                    return subprocess.CompletedProcess(
                        command, 1, "",
                        f"stage={stage}; error_number=-1708; "
                        "close_outcome=exact_document_closed_without_save; "
                        "close_failed=false; restore_failed=false",
                    )

                with patch.object(adapter.subprocess, "run", side_effect=fail):
                    with self.assertRaises(adapter.AdapterError) as failure:
                        adapter.run_request(self.request("verify_only"))
                self.assertIn(f"stage={stage};", str(failure.exception))
                self.assertIn("error_number=-1708;", str(failure.exception))
                self.assertFalse(pdf.exists())
                self.assertEqual(original, self.input.read_bytes())

    def test_private_pdf_export_transfers_identical_bytes_and_cleans_workspace(self) -> None:
        observed = {}
        target = self.root / "verification.pdf"

        def export(command):
            path = Path(command[4])
            observed["path"] = path
            self.assertNotEqual(path, target)
            self.assertEqual(Path("/private/tmp"), path.parent.parent)
            self.assertEqual(0o700, path.parent.stat().st_mode & 0o777)
            self.assertFalse(target.exists())
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(path)
            observed["bytes"] = path.read_bytes()

        with patch.object(adapter.subprocess, "run", side_effect=self.completed("verify_only", before_return=export)) as run:
            result = adapter.run_request(self.request("verify_only"))
        self.assertEqual(1, run.call_count)
        self.assertTrue(result["read_only_verified"])
        self.assertEqual(observed["bytes"], target.read_bytes())
        self.assertFalse(observed["path"].parent.exists())

    def test_private_pdf_failures_never_delete_conflicting_target(self) -> None:
        for mode in ("export", "invalid_pdf", "page_count", "conflict", "copy", "bytes", "replaced"):
            with self.subTest(mode=mode):
                target = self.root / "verification.pdf"
                target.unlink(missing_ok=True)
                observed = {}

                def produce(command):
                    path = Path(command[4])
                    observed["path"] = path
                    if mode == "invalid_pdf":
                        path.write_bytes(b"invalid")
                    else:
                        with pymupdf.open() as pdf:
                            pdf.new_page()
                            if mode == "page_count":
                                pdf.new_page()
                            pdf.save(path)
                    if mode == "conflict":
                        target.write_bytes(b"another owner")

                def copy_failure(source, destination):
                    destination.write(b"partial")
                    if mode == "replaced":
                        target.unlink()
                        target.write_bytes(b"another owner")
                    if mode in {"copy", "replaced"}:
                        raise OSError("synthetic copy failure")

                response = self.completed("verify_only", before_return=produce)
                if mode == "export":
                    def response(command, **kwargs):
                        produce(command)
                        return subprocess.CompletedProcess(command, 1, "", "synthetic export failure")
                with patch.object(adapter.subprocess, "run", side_effect=response):
                    if mode in {"copy", "bytes", "replaced"}:
                        with patch.object(adapter.shutil, "copyfileobj", side_effect=copy_failure):
                            with self.assertRaises((adapter.AdapterError, OSError)):
                                adapter.run_request(self.request("verify_only"))
                    else:
                        with self.assertRaises(adapter.AdapterError):
                            adapter.run_request(self.request("verify_only"))
                self.assertFalse(observed["path"].parent.exists())
                if mode in {"conflict", "replaced"}:
                    self.assertEqual(b"another owner", target.read_bytes())
                else:
                    self.assertFalse(target.exists())

    def test_private_pdf_timeout_removes_only_temporary_export(self) -> None:
        observed = {}
        target = self.root / "verification.pdf"

        def timeout(command, **kwargs):
            observed["path"] = Path(command[4])
            observed["path"].write_bytes(b"partial")
            target.write_bytes(b"another owner")
            raise subprocess.TimeoutExpired(command, adapter.TIMEOUT_SECONDS)

        with patch.object(adapter.subprocess, "run", side_effect=timeout):
            with self.assertRaises(adapter.AdapterError):
                adapter.run_request(self.request("verify_only"))
        self.assertEqual(b"another owner", target.read_bytes())
        self.assertFalse(observed["path"].parent.exists())

    def test_private_pdf_exclusive_create_race_preserves_other_owner(self) -> None:
        target = self.root / "verification.pdf"
        original_open = Path.open
        observed = {}

        def export(command):
            observed["path"] = Path(command[4])
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(command[4])

        def racing_open(path, mode="r", *args, **kwargs):
            if path == target and mode == "xb":
                with original_open(path, "wb") as other:
                    other.write(b"another owner")
            return original_open(path, mode, *args, **kwargs)

        with patch.object(adapter.subprocess, "run", side_effect=self.completed("verify_only", before_return=export)):
            with patch.object(Path, "open", racing_open):
                with self.assertRaises(FileExistsError):
                    adapter.run_request(self.request("verify_only"))
        self.assertEqual(b"another owner", target.read_bytes())
        self.assertFalse(observed["path"].parent.exists())

    def test_private_pdf_cleanup_error_does_not_return_successful_target(self) -> None:
        target = self.root / "verification.pdf"
        original_cleanup = tempfile.TemporaryDirectory.cleanup
        observed = {}

        def export(command):
            observed["path"] = Path(command[4])
            with pymupdf.open() as pdf:
                pdf.new_page()
                pdf.save(command[4])

        calls = []

        def fail_cleanup(workspace):
            original_cleanup(workspace)
            calls.append(workspace.name)
            if len(calls) == 1:
                raise OSError("synthetic cleanup failure")

        with patch.object(adapter.subprocess, "run", side_effect=self.completed("verify_only", before_return=export)):
            with patch.object(tempfile.TemporaryDirectory, "cleanup", fail_cleanup):
                with self.assertRaises(OSError):
                    adapter.run_request(self.request("verify_only"))
        self.assertFalse(target.exists())
        self.assertFalse(observed["path"].parent.exists())

    def test_verify_all_fields_fail_closed_without_instruction_or_result_leak(self) -> None:
        for kind in ("DATE", "SEQ", "QUOTE"):
            with self.subTest(kind=kind):
                secret_instruction = f'{kind} "private-synthetic-instruction"'
                secret_result = "private-synthetic-result"
                document = Document()
                add_field(document.add_paragraph(), " PAGE ", "1")
                add_field(document.add_paragraph(), secret_instruction, secret_result)
                document.sections[0].footer.paragraphs[0].text = ""
                document.save(self.input)
                pdf = self.root / "verification.pdf"
                pdf.unlink(missing_ok=True)

                def fail_after_pdf(command, **_kwargs):
                    Path(command[4]).write_bytes(b"partial synthetic pdf")
                    return subprocess.CompletedProcess(
                        command,
                        1,
                        "",
                        "stage=verify.all_fields.compare; error_number=7159; "
                        "close_outcome=exact_document_closed_without_save; "
                        "close_failed=false; restore_failed=false",
                    )

                with patch.object(adapter.subprocess, "run", side_effect=fail_after_pdf):
                    with self.assertRaises(adapter.AdapterError) as failure:
                        adapter.run_request(self.request("verify_only"))
                message = str(failure.exception)
                self.assertIn("verify.all_fields.compare", message)
                self.assertIn("7159", message)
                self.assertNotIn(secret_instruction, message)
                self.assertNotIn(secret_result, message)
                self.assertFalse(pdf.exists())

    def test_unsupported_field_stories_are_rejected_before_word_without_leak(self) -> None:
        original = self.input.read_bytes()
        for story in ("footnotes", "endnotes", "textbox"):
            with self.subTest(story=story):
                self.input.write_bytes(original)
                secret_instruction = f'{story.upper()} "private-synthetic-instruction"'
                secret_result = "private-synthetic-result"
                replacement = self.input.with_suffix(".unsupported.docx")
                with zipfile.ZipFile(self.input) as source, zipfile.ZipFile(
                    replacement, "w", zipfile.ZIP_DEFLATED
                ) as target:
                    for info in source.infolist():
                        data = source.read(info.filename)
                        if story == "textbox" and info.filename == "word/document.xml":
                            root = etree.fromstring(data)
                            text_box = OxmlElement("w:txbxContent")
                            paragraph = OxmlElement("w:p")
                            field = OxmlElement("w:fldSimple")
                            field.set(qn("w:instr"), secret_instruction)
                            run = OxmlElement("w:r")
                            text = OxmlElement("w:t")
                            text.text = secret_result
                            run.append(text)
                            field.append(run)
                            paragraph.append(field)
                            text_box.append(paragraph)
                            root.find(".//" + qn("w:body")).append(text_box)
                            data = etree.tostring(
                                root,
                                xml_declaration=True,
                                encoding="UTF-8",
                                standalone=True,
                            )
                        target.writestr(info, data)
                    if story in {"footnotes", "endnotes"}:
                        root = OxmlElement(f"w:{story}")
                        note = OxmlElement(
                            "w:footnote" if story == "footnotes" else "w:endnote"
                        )
                        note.set(qn("w:id"), "1")
                        paragraph = OxmlElement("w:p")
                        field = OxmlElement("w:fldSimple")
                        field.set(qn("w:instr"), secret_instruction)
                        run = OxmlElement("w:r")
                        text = OxmlElement("w:t")
                        text.text = secret_result
                        run.append(text)
                        field.append(run)
                        paragraph.append(field)
                        note.append(paragraph)
                        root.append(note)
                        target.writestr(
                            f"word/{story}.xml",
                            etree.tostring(
                                root,
                                xml_declaration=True,
                                encoding="UTF-8",
                                standalone=True,
                            ),
                        )
                replacement.replace(self.input)

                with patch.object(adapter.subprocess, "run") as run:
                    with self.assertRaisesRegex(
                        adapter.AdapterError, "does not support fields"
                    ) as failure:
                        adapter.run_request(self.request("measure_layout"))
                run.assert_not_called()
                message = str(failure.exception)
                self.assertNotIn(secret_instruction, message)
                self.assertNotIn(secret_result, message)
                self.assertNotIn(str(self.input), message)
        self.input.write_bytes(original)

    def test_input_hash_is_protected_and_refresh_output_remains_a_candidate(self) -> None:
        original = self.input.read_bytes()

        def mutate_input(_command):
            self.input.write_bytes(b"changed")

        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("measure_layout", before_return=mutate_input),
        ):
            with self.assertRaisesRegex(adapter.AdapterError, "protected input"):
                adapter.run_request(self.request("measure_layout"))
        self.input.write_bytes(original)

        def mutate_instruction(command):
            path = Path(command[3])
            replacement = path.with_suffix(".changed.docx")
            with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "word/document.xml":
                        root = etree.fromstring(data)
                        field = root.find(".//" + qn("w:fldSimple"))
                        assert field is not None
                        field.set(qn("w:instr"), " NUMPAGES ")
                        data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
                    target.writestr(info, data)
            replacement.replace(path)

        with patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed("refresh_fields", before_return=mutate_instruction),
        ):
            result = adapter.run_request(self.request("refresh_fields"))
        candidate = self.root / "refreshed.docx"
        self.assertTrue(result["field_cache_verified"])
        self.assertTrue(candidate.is_file())
        self.assertNotEqual(
            adapter._instruction_manifest(self.input, {"PAGE", "NUMPAGES"}),
            adapter._instruction_manifest(candidate, {"PAGE", "NUMPAGES"}),
        )

    def test_refresh_does_not_reparse_normal_word_manifest_reordering(self) -> None:
        before_manifest = [
            {
                "part": "word/document.xml",
                "order": 0,
                "parent_order": None,
                "form": "complex",
                "field_type": "TOC",
                "instruction": 'TOC \\o "1-3"',
            },
            {
                "part": "word/footer1.xml",
                "order": 0,
                "parent_order": None,
                "form": "complex",
                "field_type": "PAGE",
                "instruction": "PAGE",
            },
        ]
        snapshot = [
            ["synthetic-entry\t1\r"],
            [1],
            [[0, 1, 1, 1, 1, True, 1, "decimal"]],
            1,
            [["PAGE", " PAGE ", "1"]],
            [],
        ]
        response = self.response(
            "refresh_fields",
            snapshots=[snapshot, snapshot],
            toc_count=1,
            verified_count=2,
            updated_pairs=[["TOC", 1], ["PAGE", 1]],
        )
        allowed = {"TOC", "PAGE", "PAGEREF"}

        def simulate_word_normalization(command) -> None:
            candidate = Path(command[3])
            replacement = candidate.with_suffix(".normalized.docx")
            with zipfile.ZipFile(candidate) as source, zipfile.ZipFile(
                replacement, "w", zipfile.ZIP_DEFLATED
            ) as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "word/document.xml":
                        root = etree.fromstring(data)
                        original = root.find(".//" + qn("w:fldSimple"))
                        assert original is not None
                        paragraph = original.getparent()
                        insert_at = paragraph.index(original)
                        paragraph.remove(original)

                        def field_run(kind: str):
                            run = OxmlElement("w:r")
                            marker = OxmlElement("w:fldChar")
                            marker.set(qn("w:fldCharType"), kind)
                            run.append(marker)
                            return run

                        instruction_run = OxmlElement("w:r")
                        instruction = OxmlElement("w:instrText")
                        instruction.text = ' TOC \\o "1-3" '
                        instruction_run.append(instruction)
                        nested = OxmlElement("w:fldSimple")
                        nested.set(qn("w:instr"), " PAGEREF _TocSynthetic1 ")
                        nested_run = OxmlElement("w:r")
                        nested_text = OxmlElement("w:t")
                        nested_text.text = "1"
                        nested_run.append(nested_text)
                        nested.append(nested_run)
                        for offset, element in enumerate(
                            (
                                field_run("begin"),
                                instruction_run,
                                field_run("separate"),
                                nested,
                                field_run("end"),
                            )
                        ):
                            paragraph.insert(insert_at + offset, element)
                        data = etree.tostring(
                            root,
                            xml_declaration=True,
                            encoding="UTF-8",
                            standalone=True,
                        )
                    target.writestr(info, data)
                footer = OxmlElement("w:ftr")
                footer_paragraph = OxmlElement("w:p")
                footer_field = OxmlElement("w:fldSimple")
                footer_field.set(qn("w:instr"), " PAGE ")
                footer_paragraph.append(footer_field)
                footer.append(footer_paragraph)
                target.writestr(
                    "word/footer9.xml",
                    etree.tostring(
                        footer,
                        xml_declaration=True,
                        encoding="UTF-8",
                        standalone=True,
                    ),
                )
            replacement.replace(candidate)

        with patch.object(adapter, "_effective_allowed", return_value=allowed), patch.object(
            adapter, "_instruction_manifest", return_value=before_manifest
        ) as manifest, patch.object(
            adapter.subprocess,
            "run",
            side_effect=self.completed(
                "refresh_fields",
                before_return=simulate_word_normalization,
                response=response,
            ),
        ):
            result = adapter.run_request(self.request("refresh_fields"))
        self.assertEqual(1, manifest.call_count)
        self.assertTrue(result["field_cache_verified"])
        candidate = self.root / "refreshed.docx"
        self.assertTrue(candidate.is_file())
        normalized = adapter._instruction_manifest(candidate, allowed)
        self.assertTrue(
            any(
                item["field_type"] == "PAGEREF"
                and item["parent_order"] is not None
                for item in normalized
            )
        )
        self.assertTrue(
            any(item["part"] == "word/footer9.xml" for item in normalized)
        )

    def test_external_refresh_candidate_is_always_gated_by_core_writeback(self) -> None:
        response = {
            "protocol_version": "1.1",
            "status": "success",
            "operation": "refresh_fields",
            "backend": "synthetic_word",
            "software": "Microsoft Word",
            "repaginated": True,
            "saved": True,
            "field_cache_verified": True,
            "structural_changes_applied": 0,
            "updated_field_types": ["PAGE"],
        }

        def write_candidate(output: Path, mode: str) -> None:
            with zipfile.ZipFile(self.input) as source, zipfile.ZipFile(
                output, "w", zipfile.ZIP_DEFLATED
            ) as target:
                for info in source.infolist():
                    data = source.read(info.filename)
                    if info.filename == "word/document.xml" and mode != "protected_payload":
                        root = etree.fromstring(data)
                        field = root.find(".//" + qn("w:fldSimple"))
                        assert field is not None
                        if mode == "top_level_instruction":
                            field.set(qn("w:instr"), " NUMPAGES ")
                        elif mode == "field_boundary":
                            field.getparent().remove(field)
                        elif mode == "authored_body":
                            run = OxmlElement("w:r")
                            text = OxmlElement("w:t")
                            text.text = "synthetic unauthorized body change"
                            run.append(text)
                            field.getparent().append(run)
                        elif mode == "section_structure":
                            margin = root.find(".//" + qn("w:pgMar"))
                            assert margin is not None
                            margin.set(qn("w:top"), "999")
                        data = etree.tostring(
                            root,
                            xml_declaration=True,
                            encoding="UTF-8",
                            standalone=True,
                        )
                    target.writestr(info, data)
                if mode == "protected_payload":
                    target.writestr("word/media/synthetic.bin", b"protected")

        for mode in (
            "top_level_instruction",
            "field_boundary",
            "authored_body",
            "section_structure",
            "protected_payload",
        ):
            with self.subTest(mode=mode):
                candidate = self.root / f"{mode}-candidate.docx"
                final_output = self.root / f"{mode}-final.docx"

                def invoke(_command, request, _label, selected=mode):
                    write_candidate(Path(request["output_path"]), selected)
                    return subprocess.CompletedProcess(
                        [], 0, json.dumps(response), ""
                    )

                with patch.object(
                    finalize_docx, "_invoke_external_command", side_effect=invoke
                ):
                    accepted = finalize_docx.external_refresh(
                        self.input,
                        candidate,
                        json.dumps(["synthetic-adapter"]),
                        self.profile,
                        self.structure,
                        None,
                        "microsoft_word",
                    )
                self.assertEqual("synthetic_word", accepted["backend"])
                self.assertTrue(candidate.is_file())
                with patch.object(
                    finalize_docx,
                    "selective_field_result_writeback",
                    wraps=finalize_docx.selective_field_result_writeback,
                ) as writeback:
                    with self.assertRaises(finalize_docx.FormatMonographError):
                        finalize_docx.selective_field_result_writeback(
                            self.input,
                            candidate,
                            final_output,
                        )
                writeback.assert_called_once()
                self.assertFalse(final_output.exists())

    def test_path_aliases_existing_outputs_and_wrong_target_are_rejected_before_word(self) -> None:
        alias = self.root / "alias.docx"
        alias.symlink_to(self.input)
        cases = (
            self.request("measure_layout", input_path=str(alias)),
            self.request("refresh_fields", output_path=str(self.input)),
            self.request("verify_only", pdf_output_path=str(self.profile)),
            self.request("measure_layout", target_software="libreoffice"),
        )
        with patch.object(adapter.subprocess, "run") as run:
            for request in cases:
                with self.subTest(operation=request["operation"]), self.assertRaises(adapter.AdapterError):
                    adapter.run_request(request)
        run.assert_not_called()

    def test_persistent_convergence_contains_hashes_not_toc_text(self) -> None:
        toc_text = "Chapter Alpha\t1\rChapter Beta\t2\r"
        snapshot = [
            [toc_text],
            [1],
            [[0, 1, 2, 2, 1, True, 1, "lowerRoman"]],
            2,
            [["PAGE", " PAGE ", "i"]],
            [],
        ]
        sanitized = adapter._sanitize_snapshot(snapshot, 1, {"TOC", "PAGE"})
        encoded = json.dumps(sanitized)
        self.assertNotIn("Chapter Alpha", encoded)
        self.assertNotIn("Chapter Beta", encoded)
        self.assertEqual(2, sanitized["toc_entries"][0]["entry_count"])
        self.assertEqual(2, len(sanitized["toc_entries"][0]["entry_text_sha256"]))

    def test_snapshot_samples_fields_after_layout_without_new_refresh_or_retry(self) -> None:
        source = adapter.APPLESCRIPT.read_text()
        capture = source.split("on captureSnapshot(doc, allowedToken, storyPlan, stageState)", 1)[1].split("end captureSnapshot", 1)[0]
        sample = capture.index("set fieldSnapshot to my snapshotApprovedGroups")
        for observation in ("snapshot.section.start.physical", "snapshot.section.last_content", "snapshot.section.start.logical", "snapshot.total_pages.range_information", "set pageCount to pageCountRaw as integer"):
            self.assertLess(capture.index(observation), sample)
        self.assertEqual(1, capture.count("set fieldSnapshot to my snapshotApprovedGroups"))
        self.assertNotIn("updateApprovedGroups", capture)
        self.assertNotIn("delay ", capture)
        calculate = source.split("on calculateFields(doc, allowedToken, storyPlan, stageState)", 1)[1].split("end calculateFields", 1)[0]
        self.assertIn("repeat with roundIndex from 1 to 3", calculate)
        self.assertIn("my sameTuple(snapshot, item (roundIndex - 1) of snapshots)", calculate)
        verify = source.split('set currentStage of stageState to "verify.repaginate"', 1)[1].split("repeat with observedSaved", 1)[0]
        self.assertNotIn("updateApprovedGroups", verify)
        self.assertNotIn("save ownedDocument", verify)
        self.assertIn("my sameTuple(beforeAllFieldState, afterAllFieldState)", verify)

    def footer_fixture(self, *, duplicate=False, nested=False):
        doc = Document()
        doc.add_paragraph("Synthetic authored body")
        footer = doc.sections[0].footer.paragraphs[0]
        footer.text = "Page "
        add_field(footer, " PAGE ", "9")
        footer.add_run(" of ")
        add_field(footer, " NUMPAGES ", "9")
        add_field(footer, ' QUOTE "untouched" ', "unchanged")
        if duplicate:
            repeated = doc.sections[0].footer.add_paragraph("Other page context ")
            add_field(repeated, " PAGE ", "9")
        if nested:
            outer = footer._p.findall(qn("w:fldSimple"))[0]
            inner = OxmlElement("w:fldSimple")
            inner.set(qn("w:instr"), "PAGE")
            outer.append(inner)
        doc.save(self.input)
        output = self.root / "model-output.docx"
        output.write_bytes(self.input.read_bytes())
        plan = adapter._story_access_plan(self.input)
        rows = [[1, "footer_primary", "PAGE", " PAGE ", "1"],
                [1, "footer_primary", "NUMPAGES", " NUMPAGES ", "3"]]
        return output, plan, rows

    def footer_values(self, path):
        with zipfile.ZipFile(path) as package:
            root = etree.fromstring(package.read("word/footer1.xml"))
            return [(r.instruction, "".join(n.text or "" for n in adapter._result_text_nodes(root, r)))
                    for r in adapter.parse_fields(root)]

    def duplicate_footer_response(self, cache="9", values=("1", "1")):
        doc = Document()
        add_field(doc.add_paragraph(), " PAGE ", "1")
        for i in range(2):
            paragraph = doc.sections[0].footer.paragraphs[0] if i == 0 else doc.sections[0].footer.add_paragraph()
            paragraph.text = "Distinct context " + str(i) + " "
            add_field(paragraph, " PAGE ", cache)
        doc.save(self.input)
        plan = adapter._story_access_plan(self.input)
        snapshot = self.snapshot()
        snapshot[4] = [["PAGE", " PAGE ", value] for value in ("1", *values)]
        snapshot.append([self.story_observations(plan), [[1, "footer_primary", "PAGE", " PAGE ", value] for value in values]])
        return self.response("refresh_fields", [snapshot, copy.deepcopy(snapshot)],
                            verified_count=3, updated_pairs=[["PAGE", 3]],
                            story_observations=self.story_observations(plan))

    def test_refresh_rejects_unverified_duplicate_footer_caches(self):
        raw = self.duplicate_footer_response()
        before = self.input.read_bytes()
        with patch.object(adapter.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, json.dumps(raw), "")):
            with self.assertRaisesRegex(adapter.AdapterError, "footer cache"):
                adapter.run_request(self.request("refresh_fields"))
        self.assertEqual(before, self.input.read_bytes())
        self.assertFalse((self.root / "refreshed.docx").exists())

    def test_duplicate_footer_confirmation_does_not_guess_or_write(self):
        for cache, values, error in (("1", ("1", "1"), None),
                                     ("1", ("1", "2"), "cannot be matched safely")):
            with self.subTest(cache=cache, values=values):
                raw = self.duplicate_footer_response(cache, values)
                if error:
                    doc = Document(self.input)
                    doc.sections[0].footer.paragraphs[1]._p.find(".//" + qn("w:fldSimple") + "/" + qn("w:r") + "/" + qn("w:t")).text = "2"
                    doc.save(self.input)
                before = self.input.read_bytes()
                output = self.root / ("duplicate-" + values[-1] + ".docx")
                with patch.object(adapter.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, json.dumps(raw), "")):
                    if error:
                        with self.assertRaisesRegex(adapter.AdapterError, error):
                            adapter.run_request(self.request("refresh_fields", output_path=str(output)))
                        self.assertFalse(output.exists())
                    else:
                        result = adapter.run_request(self.request("refresh_fields", output_path=str(output)))
                        self.assertTrue(result["field_cache_verified"])
                        self.assertEqual(0, result["footer_cache_source"]["supplemented_fields"])
                        self.assertEqual(before, output.read_bytes())
                self.assertEqual(before, self.input.read_bytes())

    def test_footer_confirmation_is_local_and_read_only(self):
        output, plan, rows = self.footer_fixture()
        adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
        doc = Document(output)
        unrelated = doc.sections[0].footer.add_paragraph()
        add_field(unrelated, ' QUOTE "nested unrelated" ', "unchanged")
        outer = unrelated._p.find(qn("w:fldSimple"))
        nested = OxmlElement("w:fldSimple")
        nested.set(qn("w:instr"), "PAGE")
        outer.append(nested)
        doc.save(output)
        plan = adapter._story_access_plan(output)
        before = output.read_bytes()
        adapter._require_verified_footer_caches(output, {"PAGE", "NUMPAGES"}, plan, rows)
        with self.assertRaisesRegex(adapter.AdapterError, "insufficient owner-associated"):
            adapter._require_verified_footer_caches(output, {"PAGE", "NUMPAGES"}, plan, [])
        self.assertEqual(before, output.read_bytes())

    def test_footer_model_results_repair_only_authorized_cache_after_strict_validation(self):
        output, plan, rows = self.footer_fixture()
        baseline_hash = adapter.sha256(self.input)
        report = adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
        self.assertEqual(2, report["supplemented_fields"])
        self.assertEqual(baseline_hash, report["word_candidate_sha256"])
        self.assertEqual(baseline_hash, adapter.sha256(self.input))
        self.assertEqual([("PAGE", "1"), ("NUMPAGES", "3"), ('QUOTE "untouched"', "unchanged")], self.footer_values(output))
        adapter._require_saved_default_numpages(output, 3)
        adapter._require_verified_footer_caches(output, {"PAGE", "NUMPAGES"}, plan, rows)
        before = output.read_bytes()
        second = adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
        self.assertEqual(0, second["supplemented_fields"])
        self.assertEqual(before, output.read_bytes())

    def test_footer_model_results_do_not_cross_owner(self):
        output, plan, rows = self.footer_fixture()
        original = output.read_bytes()
        with self.assertRaisesRegex(adapter.AdapterError, "unapproved owner"):
            adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan,
                                             [[2, "footer_primary", "PAGE", " PAGE ", "2"]])
        self.assertEqual(original, output.read_bytes())
        doc = Document(self.input)
        section = doc.add_section(WD_SECTION.NEW_PAGE)
        section.footer.is_linked_to_previous = False
        add_field(section.footer.paragraphs[0], " PAGE ", "9")
        doc.save(self.input)
        output.write_bytes(self.input.read_bytes())
        plan = adapter._story_access_plan(self.input)
        rows.append([2, "footer_primary", "PAGE", " PAGE ", "2"])
        adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
        self.assertEqual("1", self.footer_values(output)[0][1])
        with zipfile.ZipFile(output) as package:
            root = etree.fromstring(package.read("word/footer2.xml"))
            self.assertEqual("2", "".join(root.xpath(".//w:t/text()", namespaces=adapter.NS)))

    def test_footer_saved_complex_fields_use_existing_form_and_owner_matching(self):
        from test_v032_selective_field_writeback import add_complex_field
        output, plan, rows = self.footer_fixture()
        doc = Document(output)
        footer = doc.sections[0].footer.paragraphs[0]
        footer.clear()
        footer.add_run("Page ")
        add_complex_field(footer, "PAGE", "9", dirty=False)
        footer.add_run(" of ")
        add_complex_field(footer, "NUMPAGES", "9", dirty=False)
        add_field(footer, ' QUOTE "untouched" ', "unchanged")
        doc.save(output)
        report = adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
        self.assertEqual(2, report["supplemented_fields"])
        self.assertEqual([("PAGE", "1"), ("NUMPAGES", "3"), ('QUOTE "untouched"', "unchanged")], self.footer_values(output))

    def test_footer_model_results_ambiguous_or_nested_leave_xml_route_unchanged(self):
        for duplicate, nested in ((True, False), (False, True)):
            output, plan, rows = self.footer_fixture(duplicate=duplicate, nested=nested)
            # Skipping a supplement is not evidence that saved caches are correct.
            rows = []
            before = output.read_bytes()
            report = adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
            self.assertEqual(0, report["supplemented_fields"])
            self.assertEqual(before, output.read_bytes())
            with self.assertRaisesRegex(adapter.AdapterError, "footer cache"):
                adapter._require_verified_footer_caches(output, {"PAGE", "NUMPAGES"}, plan, rows)
        output, plan, rows = self.footer_fixture(nested=True)
        before = output.read_bytes()
        self.assertEqual(0, adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)["supplemented_fields"])
        self.assertEqual(before, output.read_bytes())
        with self.assertRaisesRegex(adapter.AdapterError, "Nested footer cache"):
            adapter._require_verified_footer_caches(output, {"PAGE", "NUMPAGES"}, plan, rows)
        output, plan, rows = self.footer_fixture(duplicate=True)
        report = adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
        self.assertEqual(1, report["supplemented_fields"])
        self.assertEqual(["9", "9"], [value for instruction, value in self.footer_values(output) if instruction == "PAGE"])
        # Existing context-based XML matching remains available for duplicates.
        adapter.selective_field_result_writeback(self.input, output, self.root / "duplicate-selective.docx",
                                                  allowed_field_types={"PAGE", "NUMPAGES"})
        with self.assertRaisesRegex(adapter.AdapterError, "insufficient owner-associated"):
            adapter._require_verified_footer_caches(output, {"PAGE", "NUMPAGES"}, plan, rows)

    def test_footer_original_candidate_damage_is_rejected_before_supplement(self):
        for change in ("authored", "instruction", "boundary"):
            output, plan, rows = self.footer_fixture()
            with zipfile.ZipFile(output) as package:
                parts = {name: package.read(name) for name in package.namelist()}
            root = etree.fromstring(parts["word/footer1.xml"])
            if change == "authored":
                root.find(".//" + qn("w:t")).text = "Changed author text"
            elif change == "instruction":
                root.find(".//" + qn("w:fldSimple")).set(qn("w:instr"), "PAGE \\* ROMAN")
            else:
                field = root.find(".//" + qn("w:fldSimple"))
                field.getparent().remove(field)
            parts["word/footer1.xml"] = etree.tostring(root)
            with zipfile.ZipFile(output, "w") as package:
                for name, data in parts.items():
                    package.writestr(name, data)
            damaged = output.read_bytes()
            with self.subTest(change=change), self.assertRaises(adapter.FormatMonographError):
                adapter._supplement_footer_caches(self.input, output, {"PAGE", "NUMPAGES"}, plan, rows)
            self.assertEqual(damaged, output.read_bytes())

    def test_footer_private_associations_are_converged_and_refresh_only(self):
        output, plan, rows = self.footer_fixture()
        snapshot = self.snapshot(3)
        snapshot[4] = [row[2:] for row in rows]
        snapshot.append([self.story_observations(plan), rows])
        raw = self.response("refresh_fields", [snapshot, copy.deepcopy(snapshot)],
                            verified_count=2, updated_pairs=[["PAGE", 1], ["NUMPAGES", 1]],
                            story_observations=self.story_observations(plan))
        request = self.request("refresh_fields")
        with patch.object(adapter.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, json.dumps(raw), "")):
            result = adapter.run_request(request)
        self.assertEqual(2, result["footer_cache_source"]["supplemented_fields"])
        self.assertEqual("3", self.footer_values(Path(request["output_path"]))[1][1])
        legacy_final = copy.deepcopy(snapshot[:6])
        mixed = self.response("refresh_fields", [snapshot, legacy_final, copy.deepcopy(legacy_final)],
                              verified_count=2, updated_pairs=[["PAGE", 1], ["NUMPAGES", 1]],
                              story_observations=self.story_observations(plan))
        parsed, _ = adapter._parse_word_result(json.dumps(mixed), "refresh_fields", {"PAGE", "NUMPAGES"}, plan)
        self.assertEqual([], parsed["footer_rows"])
        raw = copy.deepcopy(raw)
        raw[13][0][6][1][0][4] = "2"
        with self.assertRaisesRegex(adapter.AdapterError, "converge"):
            adapter._parse_word_result(json.dumps(raw), "refresh_fields", {"PAGE", "NUMPAGES"}, plan)
        # Existing verify fixture has no requested footer supplement, even when
        # the helper would throw if entered.
        with patch.object(adapter, "_supplement_footer_caches", side_effect=AssertionError("verify must not supplement")) as helper, patch.object(
            adapter, "_require_verified_footer_caches", side_effect=AssertionError("refresh-only guard")
        ) as confirmation:
            verify_raw = self.response("verify_only", [snapshot, snapshot], verified_count=2,
                                       story_observations=self.story_observations(plan))
            with patch.object(adapter.subprocess, "run", return_value=subprocess.CompletedProcess([], 0, json.dumps(verify_raw), "")):
                with self.assertRaisesRegex(adapter.AdapterError, "Saved default NUMPAGES"):
                    adapter.run_request(self.request("verify_only"))
            helper.assert_not_called()
            confirmation.assert_not_called()

    def test_numpages_raw_difference_and_identically_wrong_totals_are_rejected(self) -> None:
        first = self.snapshot(page_count=3)
        first[4] = [["NUMPAGES", " NUMPAGES ", "2"]]
        second = copy.deepcopy(first)
        second[4][0][2] = "3"
        with self.assertRaisesRegex(adapter.AdapterError, "snapshots changed"):
            adapter._parse_word_result(json.dumps(self.response("verify_only", [first, second])), "verify_only", {"NUMPAGES"}, [])
        for operation in ("refresh_fields", "verify_only"):
            snapshots = [first, first]
            raw = self.response(operation, snapshots, updated_pairs=[["NUMPAGES", 1]] if operation == "refresh_fields" else [])
            with self.subTest(operation=operation), self.assertRaisesRegex(adapter.AdapterError, "NUMPAGES observation"):
                adapter._parse_word_result(json.dumps(raw), operation, {"NUMPAGES"}, [])
        parsed, _ = adapter._parse_word_result(json.dumps(self.response("verify_only", [second, second])), "verify_only", {"NUMPAGES"}, [])
        self.assertEqual(3, parsed["page_count"])
        raw = self.response("refresh_fields", [first, second, second], updated_pairs=[["NUMPAGES", 1]])
        parsed, snapshots = adapter._parse_word_result(json.dumps(raw), "refresh_fields", {"NUMPAGES"}, [])
        self.assertEqual(3, parsed["page_count"])
        self.assertEqual(3, len(snapshots))

    def test_layout_measurement_accepts_old_numpages_without_claiming_refresh(self) -> None:
        doc = Document()
        add_field(doc.add_paragraph(), " NUMPAGES ", "9")
        doc.sections[0].footer.paragraphs[0].text = ""
        doc.save(self.input)
        original = self.input.read_bytes()
        snapshot = self.snapshot(page_count=3)
        snapshot[4] = [["NUMPAGES", " NUMPAGES ", "9"]]
        valid = self.response("measure_layout", [snapshot])
        for mode in ("valid", "saved", "invalid_layout"):
            raw = copy.deepcopy(valid)
            if mode == "saved":
                raw[3] = True
            elif mode == "invalid_layout":
                raw[9] = raw[13][0][3] = 0
            with self.subTest(mode=mode), patch.object(adapter.subprocess, "run", side_effect=self.completed("measure_layout", response=raw)) as run:
                if mode == "valid":
                    response = adapter.run_request(self.request("measure_layout"))
                    self.assertEqual(3, response["page_count"])
                    self.assertTrue(response["read_only_verified"])
                    self.assertFalse(response["field_cache_verified"])
                    self.assertFalse(response["saved"])
                    self.assertEqual(adapter._hash_text("9"), response["convergence"][0]["field_results"][0]["result_sha256"])
                else:
                    with self.assertRaises(adapter.AdapterError):
                        adapter.run_request(self.request("measure_layout"))
                self.assertEqual(1, run.call_count)
            self.assertEqual(original, self.input.read_bytes())

    def test_numpages_saved_cache_gates_refresh_and_verification_publication(self) -> None:
        for operation in ("refresh_fields", "verify_only"):
            for cache in ("9", "3"):
                with self.subTest(operation=operation, cache=cache):
                    doc = Document()
                    doc.add_paragraph("Synthetic saved total")
                    footer = doc.sections[0].footer.paragraphs[0]
                    add_field(footer, " NUMPAGES ", cache)
                    # PAGE is location-dependent, not a document-wide total.
                    add_field(footer, " PAGE ", "9")
                    doc.save(self.input)
                    original = self.input.read_bytes()
                    snapshot = self.snapshot(page_count=3)
                    snapshot[4] = [["NUMPAGES", " NUMPAGES ", "3"], ["PAGE", " PAGE ", "1"]]
                    plan = adapter._story_access_plan(self.input)
                    snapshot.append([self.story_observations(plan),
                                     [[1, "footer_primary", *row] for row in snapshot[4]]])
                    raw = self.response(operation, [snapshot, snapshot], verified_count=2, updated_pairs=[["NUMPAGES", 1], ["PAGE", 1]] if operation == "refresh_fields" else [], story_observations=self.story_observations(adapter._story_access_plan(self.input)))

                    def word(command, **kwargs):
                        if operation == "verify_only":
                            with pymupdf.open() as pdf:
                                for _ in range(3):
                                    pdf.new_page()
                                pdf.save(command[4])
                        return subprocess.CompletedProcess(command, 0, json.dumps(raw), "")

                    destination = self.root / (operation + cache + (".pdf" if operation == "verify_only" else ".docx"))
                    request = self.request(operation, **{"pdf_output_path" if operation == "verify_only" else "output_path": str(destination)})
                    with patch.object(adapter.subprocess, "run", side_effect=word) as invoked:
                        if cache == "9" and operation == "verify_only":
                            with self.assertRaisesRegex(adapter.AdapterError, "Saved default NUMPAGES cache"):
                                adapter.run_request(request)
                            self.assertFalse(destination.exists())
                        else:
                            self.assertEqual("success", adapter.run_request(request)["status"])
                            self.assertTrue(destination.exists())
                            if operation == "refresh_fields":
                                self.assertEqual([("NUMPAGES", "3"), ("PAGE", "1")], self.footer_values(destination))
                        self.assertEqual(1, invoked.call_count)
                    self.assertEqual(original, self.input.read_bytes())

    def test_refresh_and_verify_final_snapshots_match_except_round(self) -> None:
        snapshot = {
            "round": 3,
            "toc_count": 1,
            "toc_entries": [
                {
                    "toc_ordinal": 1,
                    "entry_count": 1,
                    "entry_text_sha256": ["a" * 64],
                    "page_span": 1,
                }
            ],
            "sections": [
                {
                    "section_index": 0,
                    "first_physical_page": 1,
                    "last_physical_page": 1,
                    "last_content_page": 1,
                    "first_logical_page": 1,
                    "restart_numbering": True,
                    "page_number_start": 1,
                    "page_number_format": "decimal",
                }
            ],
            "page_count": 1,
            "field_results": [
                {
                    "ordinal": 1,
                    "field_type": "PAGE",
                    "instruction_sha256": "b" * 64,
                    "result_sha256": "c" * 64,
                }
            ],
            "page_boundary_spacer_ordinals": [0],
        }
        refresh = {"convergence": [copy.deepcopy(snapshot)]}
        verify_snapshot = copy.deepcopy(snapshot)
        verify_snapshot["round"] = 2
        verify = {"convergence": [verify_snapshot]}
        pdf = self.root / "matching.pdf"
        pdf.write_bytes(b"pdf")
        finalize_docx._require_matching_word_final_snapshots(refresh, verify, pdf)
        self.assertTrue(pdf.exists())

        mutations = {
            "section": lambda value: value["sections"][0].update(last_content_page=2),
            "toc_hash": lambda value: value["toc_entries"][0][
                "entry_text_sha256"
            ].__setitem__(0, "d" * 64),
            "toc_span": lambda value: value["toc_entries"][0].update(page_span=2),
            "field_instruction": lambda value: value["field_results"][0].update(
                instruction_sha256="e" * 64
            ),
            "field_result": lambda value: value["field_results"][0].update(
                result_sha256="f" * 64
            ),
            "spacer": lambda value: value.update(
                page_boundary_spacer_ordinals=[1]
            ),
            "page_count": lambda value: value.update(page_count=2),
        }
        for label, mutate in mutations.items():
            with self.subTest(label=label):
                changed = copy.deepcopy(verify_snapshot)
                mutate(changed)
                pdf = self.root / f"mismatch-{label}.pdf"
                pdf.write_bytes(b"pdf")
                with self.assertRaisesRegex(
                    finalize_docx.FormatMonographError,
                    "differs from the final field refresh snapshot",
                ):
                    finalize_docx._require_matching_word_final_snapshots(
                        refresh, {"convergence": [changed]}, pdf
                    )
                self.assertFalse(pdf.exists())

        source = (ROOT / "format-monograph" / "scripts" / "finalize_docx.py").read_text(
            encoding="utf-8"
        )
        workflow = source.split("verification = external_verify(", 1)[1]
        comparison = workflow.index("_require_matching_word_final_snapshots(")
        acceptance = workflow.index(
            'backend["read_only_verification"] = verification'
        )
        delivery = workflow.index('delivery_status = "selective_verified"')
        self.assertLess(comparison, acceptance)
        self.assertLess(acceptance, delivery)

    def test_applescript_has_no_collection_update_active_document_or_word_quit(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        lowered = source.lower()
        self.assertNotIn("active document", lowered)
        self.assertNotIn("quit application", lowered)
        self.assertNotIn("quit word", lowered)
        self.assertNotIn("update fields of", lowered)
        self.assertIn("repeat with roundIndex from 1 to 3", source)
        self.assertIn("update field fieldObject", source)
        self.assertIn("msoAutomationSecurityForceDisable", source)
        self.assertIn("set update links at open of settings to false", source)
        self.assertIn("set inputFileAlias to (POSIX file inputPath) as alias", source)
        self.assertIn("set inputPath to POSIX path of inputFileAlias", source)
        self.assertIn("set inputFileReference to inputFileAlias", source)
        self.assertIn("open inputFileReference read only false add to recent files false", source)
        self.assertIn("open file name inputPath read only true add to recent files false", source)
        self.assertNotIn("open file name inputPath read only false add to recent files false", source)
        self.assertIn("close closeTarget saving no", source)
        self.assertIn("set restoreErrors to my restoreControls", source)
        self.assertNotIn("set content of", lowered)
        self.assertNotIn("make new", lowered)

    def test_applescript_open_boundary_stages_preserve_operation_modes(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        run_body = source.split("on run argv", 1)[1]
        exact_document = source.split(
            "on exactDocument(inputPath, allowAbsent)", 1
        )[1].split("end exactDocument", 1)[0]
        self.assertEqual(1, run_body.count("with timeout of 480 seconds"))
        self.assertEqual(1, run_body.count("with timeout of"))
        alias_conversion = run_body.index(
            "set inputFileAlias to (POSIX file inputPath) as alias"
        )
        canonical_conversion = run_body.index(
            "set inputPath to POSIX path of inputFileAlias"
        )
        conversion = run_body.index("set inputFileReference to inputFileAlias")
        word_scope = run_body.index('tell application "Microsoft Word"')
        self.assertLess(alias_conversion, canonical_conversion)
        self.assertLess(canonical_conversion, conversion)
        self.assertLess(conversion, word_scope)
        self.assertEqual(2, run_body.count("inputFileReference"))
        exact_steps = (
            "copy (get posix full name of candidateDocument) to candidatePath",
            "copy candidatePath to end of candidatePaths",
            "set matchIndex to my exactPathIndex(candidatePaths, inputPath, allowAbsent)",
            "copy (get posix full name of matchedDocument) to matchedPath",
            "my exactPathIndex({matchedPath}, inputPath, false)",
        )
        for left, right in zip(exact_steps, exact_steps[1:]):
            self.assertLess(exact_document.index(left), exact_document.index(right))
        self.assertNotIn("set end of candidatePaths", exact_document)

        common_steps = (
            'set currentStage of stageState to "preopen.safe_controls.received"',
            'set currentStage of stageState to "preopen.activity_check"',
            'if (count of documents) is not 0 or (get background printing status) is not 0 then error "user_activity_conflict" number 7107',
            'set currentStage of stageState to "preopen.activity_check.received"',
            "set openAttempted to true",
            'if operationName is "refresh_fields" then',
        )
        for left, right in zip(common_steps, common_steps[1:]):
            self.assertLess(run_body.index(left), run_body.index(right))

        open_choice = run_body.split('if operationName is "refresh_fields" then', 1)[1].split(
            "end if", 1
        )[0]
        refresh_branch, readonly_branch = open_choice.split("else", 1)
        refresh_steps = (
            'set currentStage of stageState to "document.open.refresh"',
            "open inputFileReference read only false add to recent files false",
            'set currentStage of stageState to "document.open.refresh.returned"',
        )
        readonly_steps = (
            'set currentStage of stageState to "document.open.readonly"',
            "open file name inputPath read only true add to recent files false",
            'set currentStage of stageState to "document.open.readonly.returned"',
        )
        for branch_name, branch, steps in (
            ("refresh", refresh_branch, refresh_steps),
            ("readonly", readonly_branch, readonly_steps),
        ):
            with self.subTest(branch=branch_name):
                for left, right in zip(steps, steps[1:]):
                    self.assertLess(branch.index(left), branch.index(right))
        self.assertNotIn("file name inputPath", refresh_branch)
        self.assertNotIn("inputFileReference", readonly_branch)

        post_open_steps = (
            'set currentStage of stageState to "document.open.readonly.returned"',
            'set currentStage of stageState to "exact_document.lookup"',
            "set ownedDocument to my exactDocument(inputPath, false)",
            'set currentStage of stageState to "exact_document.lookup.received"',
            'set currentStage of stageState to "postopen.activity_check"',
            'if (count of documents) is not 1 or (get background printing status) is not 0 then error "user_activity_conflict" number 7107',
            'set currentStage of stageState to "postopen.activity_check.received"',
            'set currentStage of stageState to "read_only.get"',
            "set observedReadOnly to get read only of ownedDocument",
            'set currentStage of stageState to "read_only.get.received"',
            'set currentStage of stageState to "read_only.validation"',
            "if class of observedReadOnly is not boolean then error",
            'if operationName is "refresh_fields" and observedReadOnly is not false then error',
            'if operationName is not "refresh_fields" and observedReadOnly is not true then error',
            'set currentStage of stageState to "read_only.validation.received"',
        )
        for left, right in zip(post_open_steps, post_open_steps[1:]):
            self.assertLess(run_body.index(left), run_body.index(right))

        def open_contract(operation):
            if operation == "refresh_fields":
                return False, "document.open.refresh"
            if operation in {"measure_layout", "verify_only"}:
                return True, "document.open.readonly"
            raise ValueError("unsupported_operation")

        self.assertEqual((False, "document.open.refresh"), open_contract("refresh_fields"))
        self.assertEqual((True, "document.open.readonly"), open_contract("measure_layout"))
        self.assertEqual((True, "document.open.readonly"), open_contract("verify_only"))

    def test_verify_snapshots_body_and_planned_story_fields_without_persisting_values(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        group_handler = source.split(
            "on allFieldGroupState(fieldObjects, stagePrefix, stageState)", 1
        )[1].split("end allFieldGroupState", 1)[0]
        all_handler = source.split(
            "on allFieldState(doc, storyPlan, stagePrefix, stageState)", 1
        )[1].split("end allFieldState", 1)[0]
        run_body = source.split("on run argv", 1)[1]

        self.assertIn("set instructionRange to get field code of fieldObject", group_handler)
        self.assertIn("set instructionValue to get content of instructionRange", group_handler)
        self.assertIn("set fieldResultRange to get result range of fieldObject", group_handler)
        self.assertIn("set fieldResultValue to get content of fieldResultRange", group_handler)
        self.assertNotIn("kindAllowed", group_handler)
        self.assertNotIn("allowedToken", group_handler)
        self.assertIn("set bodyFields to get fields of doc", all_handler)
        self.assertIn("repeat with planReference in storyPlan", all_handler)
        self.assertNotIn("repeat with storyKind in", all_handler)
        for story in (
            "header_primary",
            "header_first",
            "header_even",
            "footer_primary",
            "footer_first",
            "footer_even",
        ):
            self.assertIn(f'"{story}"', all_handler)

        verify_steps = (
            'set currentStage of stageState to "verify.capture_snapshot.before_pdf"',
            'set currentStage of stageState to "verify.all_fields.before_pdf"',
            "set beforeAllFieldState to my allFieldState",
            'set currentStage of stageState to "verify.pdf_export"',
            "save as ownedDocument file name pdfPath file format format PDF add to recent files false",
            'set currentStage of stageState to "verify.pdf_export.reacquire"',
            "set ownedDocument to my exactDocument(inputPath, false)",
            'set currentStage of stageState to "verify.pdf_export.read_only"',
            'error "read_only_lost_after_pdf" number 7114',
            'set currentStage of stageState to "verify.pdf_export.saved"',
            'set currentStage of stageState to "verify.all_fields.after_pdf"',
            "set afterAllFieldState to my allFieldState",
            'set currentStage of stageState to "verify.all_fields.compare"',
            "set allFieldsStable to my sameTuple",
            'error "verify_all_fields_changed" number 7159',
            'set currentStage of stageState to "verify.capture_snapshot.after_pdf"',
        )
        cursor = 0
        for step in verify_steps:
            cursor = run_body.index(step, cursor) + len(step)
        returned_json = run_body.split("return my jsonValue(", 1)[1]
        self.assertNotIn("beforeAllFieldState", returned_json)
        self.assertNotIn("afterAllFieldState", returned_json)

    def test_production_snapshot_normalization_removes_only_top_level_round(self) -> None:
        snapshot = {'round': 3, 'sections': [{'round': 7}], 'page_count': 3}
        normalized = finalize_docx._final_snapshot_without_round(
            {'convergence': [snapshot]}, 'test'
        )
        self.assertEqual({'sections': [{'round': 7}], 'page_count': 3}, normalized)
        normalized['sections'][0]['round'] = 8
        self.assertEqual(7, snapshot['sections'][0]['round'])
        for invalid in ({}, {'convergence': []}, {'convergence': [None]}):
            with self.subTest(invalid=invalid), self.assertRaises(Exception):
                finalize_docx._final_snapshot_without_round(invalid, 'test')

    def test_close_stays_in_one_word_scope_and_rechecks_exact_path(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        close_handler = source.split("on closeExactDocument(inputPath)", 1)[1].split(
            "end closeExactDocument", 1
        )[0]
        exact_index = source.split(
            "on exactPathIndex(candidatePaths, inputPath, allowAbsent)", 1
        )[1].split("end exactPathIndex", 1)[0]
        run_body = source.split("on run argv", 1)[1]

        self.assertEqual(1, close_handler.count('tell application "Microsoft Word"'))
        self.assertEqual(1, close_handler.count("with timeout of 30 seconds"))
        self.assertNotIn("my exactDocument", close_handler)
        self.assertNotIn("try", close_handler)
        self.assertNotIn("on error", close_handler)
        close_steps = (
            "set documentSnapshot to get documents",
            "set candidatePaths to {}",
            "repeat with candidateDocument in documentSnapshot",
            "copy (get posix full name of candidateDocument) to end of candidatePaths",
            "set matchIndex to my exactPathIndex(candidatePaths, inputPath, true)",
            'if matchIndex is 0 then return "no_exact_document_to_close"',
            "set closeTarget to item matchIndex of documentSnapshot",
            "my exactPathIndex({get posix full name of closeTarget}, inputPath, false)",
            "close closeTarget saving no",
            "set remainingDocumentSnapshot to get documents",
            "set remainingCandidatePaths to {}",
            "repeat with remainingDocument in remainingDocumentSnapshot",
            "copy (get posix full name of remainingDocument) to end of remainingCandidatePaths",
            "set remainingMatchIndex to my exactPathIndex(remainingCandidatePaths, inputPath, true)",
            'if remainingMatchIndex is not 0 then error "exact_document_remains_after_close" number 7115',
            'return "exact_document_closed_without_save"',
        )
        for left, right in zip(close_steps, close_steps[1:]):
            self.assertLess(close_handler.index(left), close_handler.index(right))
        self.assertIn("my sameOwnedFile(candidatePath)", exact_index)
        self.assertIn('error "multiple_exact_path_matches" number 7111', exact_index)

        normal_steps = (
            'set currentStage of stageState to "document.close.normal"',
            "set closeOutcome to my closeExactDocument(inputPath)",
            'set currentStage of stageState to "document.close.normal.returned"',
        )
        cleanup_steps = (
            'set currentStage of stageState to "document.close.cleanup"',
            "set closeOutcome to my closeExactDocument(inputPath)",
            'set currentStage of stageState to "document.close.cleanup.returned"',
        )
        for steps in (normal_steps, cleanup_steps):
            start = 0
            for left, right in zip(steps, steps[1:]):
                left_at = run_body.index(left, start)
                right_at = run_body.index(right, left_at)
                self.assertLess(left_at, right_at)
                start = left_at
        self.assertEqual(2, run_body.count("my closeExactDocument(inputPath)"))
        self.assertIn(
            'set currentStage of stageState to "document.close.cleanup.failed"',
            run_body,
        )

    def test_production_identity_helpers_without_word(self) -> None:
        if adapter.sys.platform != "darwin":
            self.skipTest("System AppleScript/stat identity helper requires macOS, not Word.")
        # Execute only the production filesystem helpers, never any Word handler.
        helpers = adapter.APPLESCRIPT.read_text().split("on exactDocument(", 1)[0]
        self.assertNotIn('tell application', helpers)
        target = (self.root / "identity.docx").resolve()
        target.write_bytes(b"synthetic identity only")
        stat = target.stat()
        identity = f"{stat.st_dev}:{stat.st_ino}:Regular File"
        if not str(target).startswith("/private/var/"):
            self.skipTest("Fixture is not under the observed macOS system alias.")
        alias = "/var/" + str(target)[len("/private/var/"):]
        different = target.parent / "other" / target.name
        different.parent.mkdir()
        different.write_bytes(target.read_bytes())
        linked = target.parent / "arbitrary-link.docx"
        linked.symlink_to(target)
        hardlink = target.parent / "arbitrary-hardlink.docx"
        os.link(target, hardlink)

        def invoke(paths, allow_absent=False, completed_save_path=None):
            values = ", ".join("missing value" if value is None else json.dumps(str(value)) for value in paths)
            program = helpers + "\n" + (
                f"set ownedDocumentPath to {json.dumps(str(target))}\n"
                f"set ownedDocumentIdentity to {json.dumps(identity)}\n"
                + (f"my bindCompletedOwnedSave({json.dumps(str(completed_save_path))})\n" if completed_save_path is not None else "")
                +
                f"return my exactPathIndex({{{values}}}, ownedDocumentPath, {str(allow_absent).lower()})\n"
            )
            return subprocess.run(["/usr/bin/osascript", "-"], input=program, text=True, capture_output=True, timeout=30)

        for paths, absent, expected, error in (
            ([target], False, "1", None),
            ([alias], False, "1", None),
            ([alias], True, "1", None),  # Same rule for cleanup.
            ([], True, "0", None),
            ([different], False, None, "7112"),
            ([target, alias], False, None, "7111"),
            ([linked], False, None, "7112"),
            ([hardlink], False, None, "7112"),
            ([None], False, None, "7113"),
        ):
            with self.subTest(paths=paths, absent=absent):
                result = invoke(paths, absent)
                if error is None:
                    self.assertEqual(0, result.returncode, result.stderr)
                    self.assertEqual(expected, result.stdout.strip())
                else:
                    self.assertNotEqual(0, result.returncode)
                    self.assertIn(error, result.stderr)
        target.rename(target.with_suffix(".original"))
        missing = invoke([alias], True)
        self.assertNotEqual(0, missing.returncode)
        self.assertIn("7113", missing.stderr)
        target.write_bytes(b"different entity at the same path")
        for absent in (False, True):
            changed = invoke([alias], absent)
            self.assertNotEqual(0, changed.returncode)
            self.assertIn("7116", changed.stderr)
        # The explicit successful-save seam alone may accept normal atomic
        # replacement. A different path remains forbidden even at that seam.
        wrong_save = invoke([alias], completed_save_path=different)
        self.assertNotEqual(0, wrong_save.returncode)
        self.assertIn("7116", wrong_save.stderr)
        saved = invoke([alias], True, completed_save_path=alias)
        self.assertEqual(0, saved.returncode, saved.stderr)
        self.assertEqual("1", saved.stdout.strip())

    def test_identity_rebinding_is_only_after_approved_save_success(self) -> None:
        source = adapter.APPLESCRIPT.read_text()
        calculation = source.split("on calculateFields(", 1)[1].split("end calculateFields", 1)[0]
        order = (
            "my exactPathIndex({get posix full name of doc}, my ownedDocumentPath, false)",
            "save doc",
            'if (get saved of doc) is not true then error "calculation_copy_not_saved"',
            "my bindCompletedOwnedSave(get posix full name of doc)",
        )
        for first, second in zip(order, order[1:]):
            self.assertLess(calculation.index(first), calculation.index(second))
        self.assertEqual(1, source.count("my bindCompletedOwnedSave("))
        close = source.split("on closeExactDocument(", 1)[1].split("end closeExactDocument", 1)[0]
        self.assertNotIn("bindCompletedOwnedSave", close)

    def test_ordinary_save_is_only_in_refresh_and_does_not_claim_cache_persistence(self) -> None:
        source = adapter.APPLESCRIPT.read_text()
        calculate = source.split("on calculateFields(doc, allowedToken, storyPlan, stageState)", 1)[1].split("end calculateFields", 1)[0]
        save = "\n        save doc\n"
        self.assertEqual(1, calculate.count(save))
        self.assertNotIn("file format format document", calculate)
        self.assertNotIn("set saved of", calculate)
        self.assertNotIn("delay ", calculate)
        self.assertLess(calculate.index("if not converged then error"), calculate.index(save))
        self.assertIn("repeat with roundIndex from 1 to 3", calculate)
        verify = source.split('set currentStage of stageState to "verify.repaginate"', 1)[1].split("repeat with observedSaved", 1)[0]
        self.assertNotIn("file format format document", verify)
        self.assertNotIn("calculateFields", verify)
        self.assertIn("file format format PDF add to recent files false", verify)

    def test_applescript_uses_only_the_private_story_access_plan(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        top_level = source.split(
            "on topLevelApproved(fieldObjects, allowedToken, stageState)", 1
        )[1].split("end topLevelApproved", 1)[0]
        groups = source.split("on approvedGroups(doc, allowedToken, storyPlan, stageState)", 1)[1].split(
            "end approvedGroups", 1
        )[0]

        top_level_steps = (
            "set descriptorStarts to {}",
            "set descriptorEnds to {}",
            "set codeRange to get field code of fieldObject",
            "set resultRange to get result range of fieldObject",
            "set codeStartValue to get start of content of codeRange",
            "if class of codeStartValue is not integer then error",
            "set resultEndValue to get end of content of resultRange",
            "if class of resultEndValue is not integer then error",
            "copy codeStartValue to end of descriptorStarts",
            "copy resultEndValue to end of descriptorEnds",
        )
        for left, right in zip(top_level_steps, top_level_steps[1:]):
            self.assertLess(top_level.index(left), top_level.index(right))
        self.assertNotIn("set descriptors to", top_level)
        self.assertNotIn("set descriptor to", top_level)
        self.assertNotIn("set end of descriptors", top_level)
        self.assertIn("copy contents of candidate to end of selected", top_level)
        self.assertNotIn("set end of selected", top_level)
        selection_steps = (
            "copy item i of descriptorStarts to innerStartValue",
            "copy item i of descriptorEnds to innerEndValue",
            "copy item j of descriptorStarts to outerStartValue",
            "copy item j of descriptorEnds to outerEndValue",
            "set startInside to innerStartValue > outerStartValue",
            "set endInside to innerEndValue < outerEndValue",
            "if startInside and endInside then set nested to true",
            "set allowedKind to my kindAllowed(kindName, allowedToken)",
            "if class of allowedKind is not boolean then error",
            "set shouldSelect to false",
            "if nested is false then",
            "if allowedKind then set shouldSelect to true",
            "if shouldSelect then",
        )
        for step in selection_steps:
            with self.subTest(selection_step=step):
                self.assertIn(step, top_level)
        for unstable in (
            "set boundsReference",
            "set bounds to",
            "set boundStartReference",
            "set boundEndReference",
            "set outerBoundsReference",
            "set outerBounds to",
            "set outerStartReference",
            "set outerEndReference",
            "contents of item i of descriptor",
            "contents of item j of descriptor",
        ):
            with self.subTest(unstable=unstable):
                self.assertNotIn(unstable, top_level)
        self.assertNotIn(
            "if item 1 of bounds > item 1 of outerBounds and item 2 of bounds < item 2 of outerBounds",
            top_level,
        )
        self.assertNotIn("if not nested and my kindAllowed", top_level)

        def selected_ordinals(descriptors, kinds, allowed):
            selected = []
            for i, (bound_start, bound_end) in enumerate(descriptors):
                nested = False
                for j, (outer_start, outer_end) in enumerate(descriptors):
                    if i != j:
                        start_inside = bound_start > outer_start
                        end_inside = bound_end < outer_end
                        if start_inside and end_inside:
                            nested = True
                allowed_kind = kinds[i] in allowed
                should_select = False
                if nested is False:
                    if allowed_kind:
                        should_select = True
                if should_select:
                    selected.append(i)
            return selected

        descriptor_matrix = [
            (10, 100),
            (20, 30),
            (10, 50),
            (20, 100),
            (110, 120),
            (130, 140),
        ]
        self.assertEqual(
            [0, 2, 3, 5],
            selected_ordinals(
                descriptor_matrix,
                ["PAGE", "PAGE", "PAGE", "PAGE", "QUOTE", "REF"],
                {"PAGE", "REF"},
            ),
        )

        self.assertIn("repeat with planReference in storyPlan", groups)
        self.assertNotIn(
            'repeat with storyKind in {"header_primary", "header_first", "header_even", "footer_primary", "footer_first", "footer_even"}',
            groups,
        )
        self.assertIn('if storyObject is missing value then error "planned_story_unavailable"', groups)
        self.assertIn("if sectionIndex is 1 then", groups)
        self.assertIn("set linkedToPrevious to get link to previous of storyObject", groups)
        self.assertIn("if linkedToPrevious is true then", groups)
        self.assertIn('error "planned_story_not_independent" number 7160', groups)
        self.assertIn("set identityRows to my storyFieldIdentities", groups)
        self.assertIn("copy {sectionIndex, storyLabel, ownershipToken, storyFieldCount, identityRows} to end of storyObservations", groups)
        self.assertNotIn("storyIsOwned", groups)
        self.assertNotIn("\n    try\n", groups)
        self.assertNotIn("on error", groups)
        all_fields = source.split(
            "on allFieldState(doc, storyPlan, stagePrefix, stageState)", 1
        )[1].split("end allFieldState", 1)[0]
        self.assertIn("repeat with planReference in storyPlan", all_fields)
        self.assertNotIn("repeat with storyKind in", all_fields)
        self.assertIn('error "planned_story_unavailable" number 7160', all_fields)
        run_body = source.split("on run argv", 1)[1]
        self.assertIn("if (count of argv) is not 5 then error", run_body)
        self.assertIn("set storyPlan to my parseStoryPlan(item 5 of argv)", run_body)
        self.assertIn("savedObservations, storyObservations", run_body)

    def test_approved_field_references_are_unwrapped_before_word_commands(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        snapshot_groups = source.split(
            "on snapshotApprovedGroups(doc, groups, stageState, groupOwners)", 1
        )[1].split("end snapshotApprovedGroups", 1)[0]
        update_groups = source.split(
            "on updateApprovedGroups(groups, tocOnly, stagePrefix, stageState)", 1
        )[1].split("end updateApprovedGroups", 1)[0]
        count_groups = source.split(
            "on countApprovedFields(groups, stagePrefix, stageState)", 1
        )[1].split("end countApprovedFields", 1)[0]
        count_kinds = source.split(
            "on countKinds(groups, stagePrefix, stageState)", 1
        )[1].split("end countKinds", 1)[0]
        capture = source.split(
            "on captureSnapshot(doc, allowedToken, storyPlan, stageState)", 1
        )[1].split("end captureSnapshot", 1)[0]
        calculate = source.split(
            "on calculateFields(doc, allowedToken, storyPlan, stageState)", 1
        )[1].split("end calculateFields", 1)[0]
        run_body = source.split("on run argv", 1)[1]

        for handler_name, handler in (
            ("snapshot", snapshot_groups),
            ("update", update_groups),
            ("count", count_groups),
            ("count_kinds", count_kinds),
        ):
            with self.subTest(handler=handler_name):
                steps = (
                    "repeat with fieldGroupReference in groups",
                    "set fieldGroup to contents of fieldGroupReference",
                    "repeat with fieldReference in fieldGroup",
                    "set fieldObject to contents of fieldReference",
                )
                for left, right in zip(steps, steps[1:]):
                    self.assertLess(handler.index(left), handler.index(right))
                self.assertNotIn("repeat with fieldGroup in groups", handler)
                self.assertNotIn("repeat with fieldObject in contents of fieldGroup", handler)

        for command in (
            "set codeRange to get field code of fieldObject",
            "set resultObject to get result range of fieldObject",
        ):
            with self.subTest(snapshot_word_command=command):
                command_at = snapshot_groups.index(command)
                self.assertGreater(
                    snapshot_groups.rfind('tell application "Microsoft Word"', 0, command_at),
                    snapshot_groups.rfind("set fieldObject to contents of fieldReference", 0, command_at),
                )
                self.assertGreater(snapshot_groups.index("end tell", command_at), command_at)

        update_command = "set updateSucceeded to update field fieldObject"
        update_at = update_groups.index(update_command)
        self.assertGreater(
            update_groups.rfind('tell application "Microsoft Word"', 0, update_at),
            update_groups.rfind("set fieldObject to contents of fieldReference", 0, update_at),
        )
        self.assertGreater(update_groups.index("end tell", update_at), update_at)
        self.assertEqual(1, source.count(update_command))

        self.assertIn(
            "set fieldSnapshot to my snapshotApprovedGroups(doc, groups, stageState, item 3 of approvedResult)",
            capture,
        )
        self.assertEqual(
            2,
            calculate.count("my updateApprovedGroups(groups,"),
        )
        self.assertIn(
            'my updateApprovedGroups(groups, true, "refresh.approved_groups.toc", stageState)',
            calculate,
        )
        self.assertIn(
            'my updateApprovedGroups(groups, false, "refresh.approved_groups.field_rows", stageState)',
            calculate,
        )
        self.assertIn(
            'my countKinds(groups, "refresh.approved_groups.count", stageState)',
            calculate,
        )
        self.assertIn(
            'my countApprovedFields(groups, "verify.approved_groups", stageState)',
            run_body,
        )
        for call_site_name, call_site in (
            ("capture", capture),
            ("calculate", calculate),
            ("run", run_body),
        ):
            with self.subTest(call_site=call_site_name):
                self.assertNotIn("repeat with fieldGroup in groups", call_site)
                self.assertNotIn("count of contents of fieldGroup", call_site)

        for suffix in (
            ".repeat_setup",
            ".group_contents",
            ".field_contents",
            ".field_kind",
        ):
            with self.subTest(snapshot_stage=suffix):
                self.assertIn(f'"snapshot.approved_groups{suffix}"', snapshot_groups)
        self.assertIn('stagePrefix & ".update"', update_groups)
        self.assertIn('stagePrefix & ".count"', count_groups)
        self.assertIn('stagePrefix & ".return"', count_groups)
        self.assertIn('stagePrefix & ".count"', count_kinds)
        count_boundary_steps = (
            'set verifiedCount to my countApprovedFields(groups, "verify.approved_groups", stageState)',
            'set currentStage of stageState to "verify.approved_groups.count_received"',
            'set currentStage of stageState to "post_operation.word_version.get"',
            "set wordVersion to get version",
            'set currentStage of stageState to "post_operation.word_version.received"',
            'set currentStage of stageState to "post_operation.safe_controls"',
        )
        for left, right in zip(count_boundary_steps, count_boundary_steps[1:]):
            self.assertLess(run_body.index(left), run_body.index(right))

        source_groups = [["body-1", "body-2"], [], ["footer-1"]]
        traversed = []
        for group_reference in source_groups:
            field_group = group_reference
            for field_reference in field_group:
                field_object = field_reference
                traversed.append(field_object)
        self.assertEqual(["body-1", "body-2", "footer-1"], traversed)
        self.assertEqual([2, 0, 1], [len(group) for group in source_groups])

    def test_optional_print_preferences_have_strict_boolean_or_missing_contract(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        run_body = source.split("on run argv", 1)[1]
        assertion = source.split(
            "on assertSafePrintControls", 1
        )[1].split("end assertSafePrintControls", 1)[0]
        restore = source.split("on restoreControls", 1)[1].split(
            "end restoreControls", 1
        )[0]

        optional = (
            (
                "originalPrintFields",
                "observedPrintFields",
                "update fields at print of settings",
                "print_fields",
            ),
            (
                "originalPrintLinks",
                "observedPrintLinks",
                "update links at print of settings",
                "print_links",
            ),
            (
                "originalPrintCodes",
                "observedPrintCodes",
                "print field codes of settings",
                "print_codes",
            ),
        )
        self.assertIn("if class of originalOpenLinks is not boolean", run_body)
        self.assertIn('error "preference_original_unknown"', run_body)
        self.assertIn(
            "on assertSafePrintControls(originalPrintFields, originalPrintLinks, originalPrintCodes, stagePrefix, stageState)",
            source,
        )
        self.assertIn(
            'my assertSafePrintControls(originalPrintFields, originalPrintLinks, originalPrintCodes, "preopen.safe_controls", stageState)',
            run_body,
        )
        self.assertIn(
            'my assertSafePrintControls(originalPrintFields, originalPrintLinks, originalPrintCodes, "post_operation.safe_controls", stageState)',
            run_body,
        )
        self.assertEqual(2, run_body.count("my assertSafePrintControls("))
        for original, observed, property_name, stage_name in optional:
            with self.subTest(original=original):
                self.assertIn(f"if {original} is not missing value then set {property_name} to false", run_body)
                self.assertIn(f"if {original} is missing value then", assertion)
                self.assertIn(f"get {property_name}", assertion)
                stage_steps = (
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.get"',
                    f"set {observed} to get {property_name}",
                    f"if {original} is missing value then",
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_missing"',
                    f"if {observed} is missing value then",
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_missing.observed_missing"',
                    f"if class of {observed} is not boolean then",
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_missing.unknown"',
                    f"if {observed} is false then",
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_missing.safe_false"',
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_missing.unsafe_true"',
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_false"',
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_false.unknown"',
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_false.safe_false"',
                    f'set currentStage of stageState to stagePrefix & ".{stage_name}.expected_false.unsafe_true"',
                )
                for left, right in zip(stage_steps, stage_steps[1:]):
                    self.assertLess(assertion.index(left), assertion.index(right))
                self.assertIn(
                    f'else\n            set currentStage of stageState to stagePrefix & ".{stage_name}.expected_false"',
                    assertion,
                )
                self.assertIn(f'error "{stage_name}_state_unknown" number 7106', assertion)
                self.assertIn(f'error "{stage_name}_unsafe_true" number 7106', assertion)
                self.assertNotIn(f"if {observed} is not missing value then error", assertion)
                self.assertNotIn(f"if {observed} is not false then error", assertion)
                self.assertIn(f"if {original} is missing value then", restore)
                self.assertIn(f"set {property_name} to {original}", restore)
                missing_restore = restore.split(
                    f"if {original} is missing value then", 1
                )[1].split("else", 1)[0]
                self.assertNotIn(f"set {property_name}", missing_restore)

        # The missing-value branch is observation-only; setters occur solely in
        # the mutually exclusive Boolean branch following `else`.
        for block in (assertion,):
            self.assertNotIn("set update fields at print of settings", block)
            self.assertNotIn("set update links at print of settings", block)
            self.assertNotIn("set print field codes of settings", block)

        preopen_steps = (
            'set currentStage of stageState to "preopen.safe_controls"',
            'my assertSafePrintControls(originalPrintFields, originalPrintLinks, originalPrintCodes, "preopen.safe_controls", stageState)',
            'set currentStage of stageState to "preopen.safe_controls.received"',
        )
        post_steps = (
            'set currentStage of stageState to "post_operation.safe_controls"',
            'my assertSafePrintControls(originalPrintFields, originalPrintLinks, originalPrintCodes, "post_operation.safe_controls", stageState)',
            'set currentStage of stageState to "post_operation.safe_controls.received"',
        )
        for steps in (preopen_steps, post_steps):
            for left, right in zip(steps, steps[1:]):
                self.assertLess(run_body.index(left), run_body.index(right))

    def test_optional_print_preference_offline_state_matrix(self) -> None:
        missing = object()

        def preflight(value):
            if value is missing:
                return "missing"
            if type(value) is bool:
                return "boolean"
            raise ValueError("preference_original_unknown")

        def safe_observation(original, observed):
            if original is not missing and type(original) is not bool:
                raise ValueError("preference_original_unknown")
            if observed is missing:
                if original is missing:
                    return "observed_missing"
                raise ValueError("state_unknown")
            if type(observed) is not bool:
                raise ValueError("state_unknown")
            if observed is False:
                return "safe_false"
            raise ValueError("unsafe_true")

        def restored_observation(original, observed):
            return observed is missing if original is missing else observed is original

        for values in (
            (missing, missing, missing),
            (True, False, True),
            (missing, False, missing),
        ):
            modes = [preflight(value) for value in values]
            self.assertEqual(3, len(modes))
            safe = [missing if value is missing else False for value in values]
            restored = list(values)
            self.assertTrue(all(safe_observation(a, b) in {"observed_missing", "safe_false"} for a, b in zip(values, safe)))
            self.assertTrue(all(restored_observation(a, b) for a, b in zip(values, restored)))

        for invalid in (None, 0, "false", [], object()):
            with self.subTest(invalid=type(invalid).__name__), self.assertRaisesRegex(
                ValueError, "preference_original_unknown"
            ):
                preflight(invalid)
        for control_name in ("print_fields", "print_links", "print_codes"):
            with self.subTest(control=control_name, original="missing"):
                self.assertEqual("observed_missing", safe_observation(missing, missing))
                self.assertEqual("safe_false", safe_observation(missing, False))
                with self.assertRaisesRegex(ValueError, "unsafe_true"):
                    safe_observation(missing, True)
                for invalid in (None, 0, "false", [], object()):
                    with self.assertRaisesRegex(ValueError, "state_unknown"):
                        safe_observation(missing, invalid)
            for original in (False, True):
                with self.subTest(control=control_name, original=original):
                    self.assertEqual("safe_false", safe_observation(original, False))
                    with self.assertRaisesRegex(ValueError, "unsafe_true"):
                        safe_observation(original, True)
                    with self.assertRaisesRegex(ValueError, "state_unknown"):
                        safe_observation(original, missing)
                    for invalid in (None, 0, "false", [], object()):
                        with self.assertRaisesRegex(ValueError, "state_unknown"):
                            safe_observation(original, invalid)
        self.assertFalse(restored_observation(missing, False))
        self.assertFalse(restored_observation(True, False))

    def test_measurement_stage_labels_and_page_read_are_bounded(self) -> None:
        source = adapter.APPLESCRIPT.read_text(encoding="utf-8")
        page_at = source.split("on pageAt(", 1)[1].split("end pageAt", 1)[0]
        self.assertIn(
            "set pageRaw to get range information pointRange information type active end page number",
            page_at,
        )
        self.assertIn("if pageRaw is missing value then", page_at)
        self.assertIn("set pageValue to pageRaw as integer", page_at)
        self.assertLess(
            page_at.index("set pageRaw to get range information"),
            page_at.index("set pageValue to pageRaw as integer"),
        )
        self.assertNotIn("set pageValue to (get range information", page_at)

        labels = {
            "measure.repaginate",
            "snapshot.approved_groups",
            "approved_groups.body_fields",
            "approved_groups.body_filter",
            "approved_groups.sections",
            "approved_groups.plan_entry",
            "approved_groups.section_object",
            "approved_groups.story_loop_complete",
            "approved_groups.section_loop_complete",
            "approved_groups.return_prepare",
            "approved_groups.return_prepare.group",
            "approved_groups.return_prepare.field",
            "approved_groups.return_execute",
            "snapshot.approved_groups.received",
            "snapshot.approved_groups.first_use",
            "refresh.approved_groups.toc.received",
            "refresh.approved_groups.toc.first_use",
            "refresh.approved_groups.field_rows.received",
            "refresh.approved_groups.field_rows.first_use",
            "refresh.approved_groups.count.received",
            "refresh.approved_groups.count.first_use",
            "verify.approved_groups.received",
            "verify.approved_groups.first_use",
            "snapshot.toc.result",
            "snapshot.toc.page_span.start",
            "snapshot.toc.page_span.end",
            "snapshot.section.start",
            "snapshot.section.end",
            "snapshot.section.last_content",
            "snapshot.section.start.logical",
            "snapshot.section.page_number_options",
            "snapshot.total_pages.range_information",
            "snapshot.field_rows.code",
            "snapshot.field_rows.result",
            "snapshot.spacer_rows",
        }
        for label in labels:
            with self.subTest(label=label):
                self.assertIn(f'"{label}"', source)
        self.assertEqual(5, source.count("my approvedGroups("))
        for base_label in (
            "snapshot.approved_groups",
            "refresh.approved_groups.toc",
            "refresh.approved_groups.field_rows",
            "refresh.approved_groups.count",
            "verify.approved_groups",
        ):
            with self.subTest(call_stage=base_label):
                call_stage = source.index(
                    f'set currentStage of stageState to "{base_label}"'
                )
                call = source.index("my approvedGroups(", call_stage)
                received = source.index(
                    f'set currentStage of stageState to "{base_label}.received"', call
                )
                first_use = source.index(
                    f'set currentStage of stageState to "{base_label}.first_use"',
                    received,
                )
                self.assertLess(call_stage, call)
                self.assertLess(call, received)
                self.assertLess(received, first_use)
        top_level = source.split(
            "on topLevelApproved(fieldObjects, allowedToken, stageState)", 1
        )[1].split("end topLevelApproved", 1)[0]
        for suffix in (
            ".descriptor_start_append",
            ".descriptor_end_append",
            ".selection_setup",
            ".selection_inner_start",
            ".selection_inner_end",
            ".selection_outer_start",
            ".selection_outer_end",
            ".selection_compare_start",
            ".selection_compare_end",
            ".selection_mark_nested",
            ".selection_candidate",
            ".selection_kind",
            ".selection_allowed_kind",
            ".selection_decision",
            ".selection_append",
            ".return",
        ):
            with self.subTest(selection_suffix=suffix):
                self.assertIn(f'"{suffix}"', top_level)
        groups = source.split("on approvedGroups(doc, allowedToken, storyPlan, stageState)", 1)[1].split(
            "end approvedGroups", 1
        )[0]
        self.assertIn("repeat with planReference in storyPlan", groups)
        self.assertNotIn("repeat with storyKind in", groups)
        for suffix in (
            ".story_object",
            ".ownership_first_section",
            ".link_state",
            ".link_state_validation",
            ".ownership_independent",
            ".branch_complete",
            ".story_text_object",
            ".story_fields",
            ".story_filter",
            ".append",
        ):
            with self.subTest(suffix=suffix):
                self.assertIn(f'"{suffix}"', groups)
        self.assertIn(
            "return {tocTexts, tocSpans, sectionRows, pageCount, fieldRows, spacerRows, {item 2 of approvedResult, item 4 of fieldSnapshot}}",
            source,
        )

        run_body = source.split("on run argv", 1)[1]
        failure = run_body.split(
            "on error failureMessage number failureNumber", 1
        )[1]
        wrapper = next(
            line.strip()
            for line in failure.splitlines()
            if line.strip().startswith('error ("stage="')
        )
        self.assertIn("failureStage", wrapper)
        for forbidden in (
            "failureMessage",
            "inputPath",
            "pdfPath",
            "codeText",
            "resultText",
            "tocTexts",
        ):
            self.assertNotIn(forbidden, wrapper)
        self.assertNotIn("log ", source)


if __name__ == "__main__":
    unittest.main()
