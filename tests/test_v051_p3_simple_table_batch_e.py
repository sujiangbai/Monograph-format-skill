"""P3-E public, synthetic, offline CLI and deterministic-format checks."""
from __future__ import annotations

import copy
import contextlib
import io
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'format-monograph/scripts'))
import structure_map as sm


class SimpleTableBatchETests(unittest.TestCase):
    def fixture(self, folder, *, caption=True, rows=3, mutate=None, combined=False):
        document = Document()
        document.add_paragraph('Synthetic ordinary body')
        if caption:
            document.add_paragraph('表 1-1 Existing synthetic caption\nSecond authored line', 'Caption')
        table = document.add_table(rows=rows, cols=2)
        table.autofit = False
        for column in table.columns:
            column.width = Inches(1)
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.width = Inches(1)
                cell.text = f'Item {row_index}' if column_index == 0 else str(row_index)
                cell.paragraphs[0].runs[0].italic = True
        document.add_paragraph('Existing table note: unchanged', 'Caption')
        other = document.add_table(rows=2, cols=2)
        other.cell(0, 0).text = 'Non-target table'
        document.sections[0].header.paragraphs[0].text = 'Unchanged header'
        if combined:
            from test_v051_p3_image_caption_batch_d import PNG
            document.add_paragraph().add_run().add_picture(io.BytesIO(PNG), width=Inches(1))
            document.add_paragraph('图 1-1 Existing synthetic figure caption', 'Caption')
            document.add_paragraph('Synthetic following body')
        if mutate:
            mutate(document)
        source = folder / 'synthetic.docx'
        document.save(source)
        mapping = sm.candidate_structure_map(source)
        mapping['status'] = 'approved'
        mapping['tables'][0].update(approved=True, kind='data', header_rows=[0],
                                    repeat_header_rows=[0], prevent_normal_row_split=True)
        mapping['tables'][0]['visual'] = {
            'approved': True, 'alignment': 'center', 'text_wrapping': 'none',
            'border_preset': 'three_line', 'all_cell_alignment': 'center',
            'column_roles': ['unit', 'numeric'],
        }
        for entry in mapping['paragraph_roles']:
            if entry['locator'].get('paragraph') == 0:
                entry.update(approved=True, role='body_text', canonical_role='body_text')
            if caption and entry['locator'].get('paragraph') == 1:
                entry.update(approved=True, role='table_caption', canonical_role='table_caption')
        if caption:
            # A caller can approve an existing multiline caption which the old
            # candidate heuristic does not recognize. Do not change that heuristic.
            locator = copy.deepcopy(mapping['paragraph_roles'][1]['locator'])
            mapping['captions'] = [e for e in mapping['captions'] if e.get('locator') != locator] + [{
                'locator': locator, 'text_sha256': sm.text_sha256(document.paragraphs[1].text),
                'label': '表', 'sequence_name': 'Table', 'numbering_mode': 'manual_text',
                'identifier_semantics': 'publication_number', 'domain_context': 'general',
                'domain_confidence': 'high', 'approved': True, 'action': 'style_only',
            }]
            mapping['pagination_groups'] = [g for g in mapping['pagination_groups'] if g.get('kind') != 'table_caption_with_table'] + [{
                'kind': 'table_caption_with_table', 'anchor': locator, 'table': 0,
                'table_text_sha256': mapping['tables'][0]['table_text_sha256'], 'approved': True,
            }]
            for entry in mapping['captions']:
                if entry['locator'].get('paragraph') == 1:
                    entry.update(approved=True, action='style_only')
            for group in mapping['pagination_groups']:
                if group['kind'] == 'table_caption_with_table' and group['table'] == 0:
                    group['approved'] = True
        if combined:
            for group in mapping['pagination_groups']:
                if group['kind'] == 'figure_with_caption':
                    group['approved'] = True
            for entry in mapping['captions']:
                if entry.get('sequence_name') == 'Figure':
                    entry.update(approved=True, action='style_only')
            for entry in mapping['paragraph_roles']:
                index = entry['locator'].get('paragraph')
                if index == 4:
                    entry.update(approved=True, role='figure_caption', canonical_role='figure_caption')
                if index == 5:
                    entry.update(approved=True, role='body_text', canonical_role='body_text')
        profile = json.loads((ROOT / 'format-monograph/examples/profiles/v051-foundation-format-slice.json').read_text())
        mp, pp = folder / 'map.json', folder / 'profile.json'
        mp.write_text(json.dumps(mapping))
        pp.write_text(json.dumps(profile))
        return source, mp, pp, mapping

    def cli(self, script, *args):
        return subprocess.run([sys.executable, '-B', str(ROOT / 'format-monograph/scripts' / script),
                               *map(str, args)], capture_output=True, text=True, timeout=60)

    def apply_cli(self, source, mp, pp, output):
        result = self.cli('apply_profile.py', source, '--structure-map', mp, '--profile', pp,
                          '--output-dir', output, '--allow-missing-fonts')
        self.assertEqual(0, result.returncode, result.stdout + result.stderr)
        return output / (source.stem + '-formatted.docx')

    def test_real_cli_supported_table_and_non_targets(self):
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            root = Path(folder)
            source, mp, pp, mapping = self.fixture(root)
            original = source.read_bytes()
            output = self.apply_cli(source, mp, pp, root / 'out')
            before, after = Document(source), Document(output)
            self.assertEqual(WD_TABLE_ALIGNMENT.CENTER, after.tables[0].alignment)
            paragraph = after.tables[0].cell(0, 0).paragraphs[0]
            self.assertEqual(WD_LINE_SPACING.AT_LEAST, paragraph.paragraph_format.line_spacing_rule)
            self.assertEqual(Pt(15), paragraph.paragraph_format.line_spacing)
            self.assertEqual(Pt(9), paragraph.runs[0].font.size)
            self.assertTrue(paragraph.runs[0].italic)
            self.assertEqual(before.tables[1]._tbl.xml, after.tables[1]._tbl.xml)
            self.assertEqual(before.paragraphs[-1]._p.xml, after.paragraphs[-1]._p.xml)
            self.assertEqual(before.styles['Caption'].element.xml, after.styles['Caption'].element.xml)
            self.assertEqual(before.tables[0]._tbl.tblGrid.xml, after.tables[0]._tbl.tblGrid.xml)
            self.assertEqual(before.tables[0].autofit, after.tables[0].autofit)
            self.assertEqual(original, source.read_bytes())
            self.assertTrue(after.paragraphs[0].style.name.startswith('Monograph Approved '))
            self.assertEqual(Pt(6), after.paragraphs[1].paragraph_format.space_before)
            self.assertEqual(Pt(3), after.paragraphs[1].paragraph_format.space_after)
            self.assertTrue(after.paragraphs[1].paragraph_format.keep_with_next)
            self.assertEqual(before.paragraphs[1].text, after.paragraphs[1].text)
            audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
            self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)

    def rule(self):
        return {'id': sm.FOUNDATION_TABLE_RULE, 'properties': copy.deepcopy(sm.FOUNDATION_TABLE_PROPERTIES)}

    def test_false_row_flags_are_enabled_and_audit_rejects_false_cli(self):
        for value in ('0', 'false', 'off'):
            with self.subTest(value=value), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                def mutate(doc):
                    for table in doc.tables:
                        for row in table.rows:
                            for tag in ('tblHeader', 'cantSplit'):
                                node = OxmlElement('w:' + tag); node.set(qn('w:val'), value)
                                row._tr.get_or_add_trPr().append(node)
                source, mp, pp, _ = self.fixture(root, mutate=mutate)
                output = self.apply_cli(source, mp, pp, root / 'out')
                before, after = Document(source), Document(output)
                self.assertEqual(before.tables[1]._tbl.xml, after.tables[1]._tbl.xml)
                for index, row in enumerate(after.tables[0].rows):
                    self.assertEqual(['true'], [n.get(qn('w:val')) for n in row._tr.xpath('./w:trPr/w:cantSplit')])
                    self.assertEqual(['true'] if index == 0 else [],
                                     [n.get(qn('w:val')) for n in row._tr.xpath('./w:trPr/w:tblHeader')])
                audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)
                for tag in ('tblHeader', 'cantSplit'):
                    changed = Document(output)
                    changed.tables[0].rows[0]._tr.xpath('./w:trPr/w:' + tag)[0].set(qn('w:val'), value)
                    bad = root / (tag + '.docx'); changed.save(bad)
                    audit = self.cli('audit_docx.py', source, bad, '--structure-map', mp, '--profile', pp)
                    self.assertNotEqual(0, audit.returncode)

    def test_conflicting_visual_requests_preserved_and_reported_cli(self):
        cases = ({'header_shading_hex': 'FF0000'}, {'horizontal_rule_rows': [1]},
                 {'inside_vertical_borders': True}, {'major_border_pt': 2},
                 {'column_roles': {'unit': True, 'numeric': True}},
                 {'landscape_approved': 'false'}, {'cell_margins_mm': {'top': True}})
        for request in cases:
            with self.subTest(request=request), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                source, mp, pp, mapping = self.fixture(root)
                mapping['paragraph_roles'][0]['approved'] = False
                mapping['tables'][0]['visual'].update(request); mp.write_text(json.dumps(mapping))
                profile = json.loads(pp.read_text()); profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_TABLE_RULE]
                pp.write_text(json.dumps(profile))
                # Loader rejection is also acceptable, provided no output is published.
                result = self.cli('apply_profile.py', source, '--structure-map', mp, '--profile', pp,
                                  '--output-dir', root / 'out', '--allow-missing-fonts')
                output = root / 'out/synthetic-formatted.docx'
                if result.returncode:
                    self.assertFalse(output.exists())
                    self.assertNotIn('Profile validation failed', result.stderr)
                    self.assertTrue('P3E' in result.stderr or 'Structure map' in result.stderr, result.stderr)
                    continue
                before, after = Document(source), Document(output)
                self.assertEqual(before.element.xml, after.element.xml)
                self.assertEqual(before.styles.element.xml, after.styles.element.xml)
                audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)
                self.assertIn('unsupported_or_conflicting_visual_request', audit.stdout)
                self.assertIn('"tables": 0', audit.stdout)
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            source, mp, pp, mapping = self.fixture(Path(folder))
            mapping['tables'][0]['visual'].update(vertical_alignment='center', orientation='portrait',
                                                  landscape_approved=False,
                                                  cell_margins_mm={'top': 1.0, 'bottom': 1.0, 'left': 1.5, 'right': 1.5})
            self.assertEqual(1, len(sm._foundation_table_plan(Document(source), mapping, self.rule())[0]))

    def test_explicit_oversized_rows_skip_actual_section_without_mutation_cli(self):
        from docx.enum.section import WD_SECTION_START
        from docx.enum.table import WD_ROW_HEIGHT_RULE
        for height_rule in (WD_ROW_HEIGHT_RULE.EXACTLY, WD_ROW_HEIGHT_RULE.AT_LEAST):
            with self.subTest(height_rule=height_rule), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                def mutate(doc):
                    section = doc.sections[0]
                    section.page_height = Inches(5); section.top_margin = section.bottom_margin = Inches(1)
                    doc.tables[0].rows[1].height = Inches(4)
                    doc.tables[0].rows[1].height_rule = height_rule
                    last = doc.add_section(WD_SECTION_START.NEW_PAGE); last.page_height = Inches(12)
                source, mp, pp, mapping = self.fixture(root, mutate=mutate)
                mapping['paragraph_roles'][0]['approved'] = False
                mp.write_text(json.dumps(mapping))
                profile = json.loads(pp.read_text()); profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_TABLE_RULE]; pp.write_text(json.dumps(profile))
                original = source.read_bytes()
                plans, skipped = sm._foundation_table_plan(Document(source), mapping, self.rule())
                self.assertEqual([], plans)
                self.assertIn({'table': 0, 'rows': [1], 'reason': 'explicit_row_exceeds_section_height'}, skipped)
                output = self.apply_cli(source, mp, pp, root / 'out')
                before, after = Document(source), Document(output)
                self.assertEqual(before.element.xml, after.element.xml)
                self.assertEqual(before.styles.element.xml, after.styles.element.xml)
                self.assertEqual(original, source.read_bytes())
                self.assertFalse(after.tables[0].rows[1]._tr.xpath('./w:trPr/w:cantSplit'))
                audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)
                self.assertIn('explicit_row_exceeds_section_height', audit.stdout)
                self.assertIn('"tables": 0', audit.stdout)

    def test_explicit_no_cell_shading_overrides_direct_theme_and_shared_style_cli(self):
        from docx.enum.style import WD_STYLE_TYPE
        for kind in ('direct', 'theme', 'inherited_conditional'):
            with self.subTest(kind=kind), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                def shade(parent, theme=False):
                    node = OxmlElement('w:shd')
                    node.set(qn('w:val'), 'pct20')
                    node.set(qn('w:fill'), 'FF0000')
                    if theme:
                        for key, value in {'themeFill': 'accent1', 'themeFillTint': '80',
                                           'themeColor': 'accent2', 'themeShade': '80'}.items():
                            node.set(qn('w:' + key), value)
                    parent.append(node)
                def mutate(doc):
                    style = doc.styles.add_style('Synthetic shared shading', WD_STYLE_TYPE.TABLE)
                    props = OxmlElement('w:tcPr'); style.element.append(props); shade(props)
                    conditional = OxmlElement('w:tblStylePr')
                    conditional.set(qn('w:type'), 'firstRow'); style.element.append(conditional)
                    props = OxmlElement('w:tcPr'); conditional.append(props); shade(props, True)
                    for table in doc.tables:
                        table.style = style
                    cell = doc.tables[0].cell(0, 0)
                    if kind != 'inherited_conditional':
                        shade(cell._tc.get_or_add_tcPr(), kind == 'theme')
                    p = cell.paragraphs[0]
                    shade(p._p.get_or_add_pPr())
                    shade(p.runs[0]._r.get_or_add_rPr())
                    highlight = OxmlElement('w:highlight'); highlight.set(qn('w:val'), 'yellow')
                    p.runs[0]._r.get_or_add_rPr().append(highlight)
                source, mp, pp, mapping = self.fixture(root, mutate=mutate)
                profile = json.loads(pp.read_text())
                profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_TABLE_RULE]
                pp.write_text(json.dumps(profile))
                before = Document(source)
                first = self.apply_cli(source, mp, pp, root / 'first')
                after = Document(first)
                for row in after.tables[0].rows:
                    for cell in row.cells:
                        nodes = cell._tc.xpath('./w:tcPr/w:shd')
                        self.assertEqual(1, len(nodes))
                        self.assertEqual({qn('w:val'): 'nil', qn('w:color'): 'auto',
                                          qn('w:fill'): 'auto'}, dict(nodes[0].attrib))
                self.assertEqual(before.styles.element.xml, after.styles.element.xml)
                self.assertEqual(before.tables[1]._tbl.xml, after.tables[1]._tbl.xml)
                query = './/w:pPr/w:shd | .//w:rPr/w:shd | .//w:rPr/w:highlight'
                self.assertEqual([etree.tostring(n, method='c14n') for n in before.tables[0]._tbl.xpath(query)],
                                 [etree.tostring(n, method='c14n') for n in after.tables[0]._tbl.xpath(query)])
                audit = self.cli('audit_docx.py', source, first, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)
                second = self.apply_cli(first, mp, pp, root / 'second')
                self.assertEqual(after.element.xml, Document(second).element.xml)
                for corruption in ('color', 'theme', 'inherit'):
                    changed = Document(first)
                    node = changed.tables[0].cell(0, 0)._tc.xpath('./w:tcPr/w:shd')[0]
                    if corruption == 'inherit':
                        node.getparent().remove(node)
                    else:
                        node.set(qn('w:val'), 'clear')
                        node.set(qn('w:themeFill' if corruption == 'theme' else 'w:fill'),
                                 'accent1' if corruption == 'theme' else 'FF0000')
                    bad = root / (corruption + '.docx'); changed.save(bad)
                    audit = self.cli('audit_docx.py', source, bad, '--structure-map', mp, '--profile', pp)
                    self.assertNotEqual(0, audit.returncode, corruption)

    def test_no_caption_and_long_table_remain_structural_not_visual_proof(self):
        for rows in (3, 55):
            with self.subTest(rows=rows), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                source, mp, pp, mapping = self.fixture(root, caption=False, rows=rows)
                output = self.apply_cli(source, mp, pp, root / 'out')
                original, result = Document(source), Document(output)
                self.assertEqual([p.text for p in original.paragraphs], [p.text for p in result.paragraphs])
                self.assertEqual(rows, len(result.tables[0].rows))
                for index, row in enumerate(result.tables[0].rows):
                    self.assertEqual(index == 0, bool(row._tr.xpath('./w:trPr/w:tblHeader')))
                    self.assertTrue(row._tr.xpath('./w:trPr/w:cantSplit'))
                report = (root / 'out/synthetic-format-report.md').read_text()
                self.assertIn('pending_visual_verification', report)
                self.assertIn('"captions": 0', report)

    def test_exact_borders_margins_and_character_properties(self):
        def overrides(document):
            table = document.tables[0]
            table.style = 'Table Grid'
            for row in table.rows:
                for cell in row.cells:
                    tc_pr = cell._tc.get_or_add_tcPr()
                    borders = OxmlElement('w:tcBorders')
                    for side in ('top', 'bottom', 'left', 'right', 'tl2br'):
                        border = OxmlElement('w:' + side)
                        border.set(qn('w:val'), 'double')
                        borders.append(border)
                    tc_pr.append(borders)
                    tc_pr.append(OxmlElement('w:tcMar'))
                    p = cell.paragraphs[0]
                    p.paragraph_format.first_line_indent = Pt(-9)
                    p.paragraph_format.space_before = Pt(12)
                    p.runs[0].bold = True
                    p.runs[0].underline = True
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(root, mutate=overrides)
            output = self.apply_cli(source, mp, pp, root / 'out')
            table = Document(output).tables[0]
            for index, row in enumerate(table.rows):
                for cell in row.cells:
                    borders = cell._tc.tcPr.find(qn('w:tcBorders'))
                    for side in ('top', 'bottom', 'left', 'right', 'insideH', 'insideV', 'tl2br'):
                        border = borders.find(qn('w:' + side))
                        expected = 8 if (side == 'top' and index == 0) or (side == 'bottom' and index == 2) else 4 if side == 'bottom' and index == 0 else 0
                        self.assertEqual(str(expected), border.get(qn('w:sz')))
                        self.assertEqual('single' if expected else 'nil', border.get(qn('w:val')))
                    p = cell.paragraphs[0]
                    self.assertEqual(Pt(0), p.paragraph_format.first_line_indent)
                    self.assertEqual(Pt(0), p.paragraph_format.space_before)
                    self.assertTrue(p.runs[0].bold and p.runs[0].italic and p.runs[0].underline)
                    self.assertFalse(cell._tc.xpath('./w:tcPr/w:tcMar'))
                    fonts = p.runs[0]._r.rPr.rFonts
                    self.assertEqual('Songti', fonts.get(qn('w:eastAsia')))
                    self.assertEqual('Times New Roman', fonts.get(qn('w:ascii')))
            margins = table._tbl.tblPr.find(qn('w:tblCellMar'))
            self.assertEqual({'top': '57', 'bottom': '57', 'left': '85', 'right': '85'},
                             {e.tag.split('}')[1]: e.get(qn('w:w')) for e in margins})

    def test_unapproved_missing_header_and_special_column_preserved_cli(self):
        for case in ('unapproved', 'visual_unapproved', 'no_header', 'multi_header', 'special_column', 'width_request'):
            with self.subTest(case=case), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                source, mp, pp, mapping = self.fixture(root)
                entry = mapping['tables'][0]
                if case == 'unapproved': entry['approved'] = False
                if case == 'visual_unapproved': entry['visual']['approved'] = False
                if case == 'no_header': entry['header_rows'] = []
                if case == 'multi_header': entry['header_rows'] = [0, 1]
                if case == 'special_column': entry['visual']['column_roles'] = ['short_code', 'numeric']
                if case == 'width_request': entry['visual']['available_width_percent'] = 100
                mp.write_text(json.dumps(mapping))
                output = self.apply_cli(source, mp, pp, root / 'out')
                a, b = Document(source), Document(output)
                self.assertEqual(a.tables[0]._tbl.xml, b.tables[0]._tbl.xml)
                self.assertEqual(a.paragraphs[1]._p.xml, b.paragraphs[1]._p.xml)
                audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)
                self.assertIn('not_applicable', audit.stdout)
                self.assertIn('"tables": 0', audit.stdout)

    def test_complex_topology_and_payload_zero_write(self):
        for case in ('merged', 'nested', 'field', 'hyperlink', 'floating', 'omml', 'image', 'list', 'row_container'):
            with self.subTest(case=case), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                def mutate(document):
                    table = document.tables[0]
                    if case == 'merged': table.cell(0, 0).merge(table.cell(0, 1))
                    if case == 'nested': table.cell(1, 0).add_table(rows=1, cols=1)
                    if case == 'floating': table._tbl.tblPr.append(OxmlElement('w:tblpPr'))
                    if case == 'row_container':
                        control, content = OxmlElement('w:sdt'), OxmlElement('w:sdtContent')
                        content.append(table.rows[1]._tr)
                        control.append(content)
                        table._tbl.append(control)
                    if case in ('field', 'hyperlink', 'omml', 'image'):
                        tag = {'field': 'w:fldSimple', 'hyperlink': 'w:hyperlink', 'omml': 'm:oMath', 'image': 'w:drawing'}[case]
                        table.cell(1, 0).paragraphs[0]._p.append(OxmlElement(tag))
                    if case == 'list': table.cell(1, 0).paragraphs[0]._p.get_or_add_pPr().append(OxmlElement('w:numPr'))
                source, _, _, mapping = self.fixture(root, mutate=mutate)
                document = Document(source)
                before = document.element.xml
                evidence = sm._apply_foundation_simple_tables(document, mapping, self.rule())
                self.assertEqual(0, evidence['tables'])
                self.assertTrue(evidence['skipped'])
                self.assertEqual(before, document.element.xml)

    def test_identity_and_legacy_operation_conflicts_fail_before_write(self):
        for case in ('hash', 'row_count', 'duplicate', 'pair_hash', 'duplicate_pair', 'cleanup', 'move', 'properties'):
            with self.subTest(case=case), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                source, _, _, mapping = self.fixture(Path(folder))
                entry = mapping['tables'][0]
                rule = self.rule()
                if case == 'hash': entry['table_text_sha256'] = '0' * 64
                if case == 'row_count': entry['row_count'] = 99
                if case == 'duplicate': mapping['tables'].append(copy.deepcopy(entry))
                if case == 'pair_hash': mapping['pagination_groups'][0]['table_text_sha256'] = '0' * 64
                if case == 'duplicate_pair': mapping['pagination_groups'] *= 2
                if case == 'cleanup': mapping['table_cell_cleanups'] = [{'approved': True}]
                if case == 'move': mapping['captions'][0]['action'] = 'move_caption'
                if case == 'properties': rule['properties']['font_size_pt'] = 12
                document = Document(source)
                before = document.element.xml
                with self.assertRaises(sm.FormatMonographError):
                    sm._apply_foundation_simple_tables(document, mapping, rule)
                self.assertEqual(before, document.element.xml)

    def test_caption_wrong_position_or_missing_approval_never_moves(self):
        for case in ('below', 'gap', 'role', 'action', 'duplicate_role'):
            with self.subTest(case=case), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                source, _, _, mapping = self.fixture(Path(folder))
                document = Document(source)
                if case == 'below': document.tables[0]._tbl.addnext(document.paragraphs[1]._p)
                if case == 'gap': document.paragraphs[1]._p.addnext(OxmlElement('w:p'))
                if case == 'role': mapping['paragraph_roles'][1]['approved'] = False
                if case == 'action': mapping['captions'][0]['approved'] = False
                if case == 'duplicate_role': mapping['paragraph_roles'].append(copy.deepcopy(mapping['paragraph_roles'][1]))
                before = document.element.xml
                evidence = sm._apply_foundation_simple_tables(document, mapping, self.rule())
                self.assertEqual(0, evidence['tables'])
                self.assertEqual(before, document.element.xml)

    def test_audit_rejects_non_whitelisted_table_and_caption_changes(self):
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(root)
            output = self.apply_cli(source, mp, pp, root / 'out')
            saved = output.read_bytes()
            for case in ('text', 'bold', 'width', 'grid', 'height', 'border', 'margin', 'header', 'font', 'spacing', 'position', 'other_table', 'caption', 'shared_style'):
                with self.subTest(case=case):
                    document = Document(io.BytesIO(saved))
                    table = document.tables[0]
                    paragraph = table.cell(1, 0).paragraphs[0]
                    if case == 'text': paragraph.add_run('changed')
                    if case == 'bold': paragraph.runs[0].bold = True
                    if case == 'width': table.cell(1, 0).width = Inches(3)
                    if case == 'grid': table.columns[0].width = Inches(3)
                    if case == 'height': table.rows[1].height = Inches(3)
                    if case == 'border': table._tbl.tblPr.find(qn('w:tblBorders'))[0].set(qn('w:sz'), '20')
                    if case == 'margin': table._tbl.tblPr.find(qn('w:tblCellMar'))[0].set(qn('w:w'), '500')
                    if case == 'header': table.rows[1]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                    if case == 'font': paragraph.runs[0].font.size = Pt(12)
                    if case == 'spacing': paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    if case == 'position': document.element.body.insert(0, table._tbl)
                    if case == 'other_table': document.tables[1].alignment = WD_TABLE_ALIGNMENT.RIGHT
                    if case == 'caption': document.paragraphs[1].paragraph_format.keep_with_next = False
                    if case == 'shared_style': document.styles['Caption'].font.size = Pt(18)
                    document.save(output)
                    audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                    self.assertNotEqual(0, audit.returncode, case + audit.stdout)

    def test_same_approval_repeat_is_semantically_idempotent_cli(self):
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(root)
            first = self.apply_cli(source, mp, pp, root / 'first')
            second = self.apply_cli(first, mp, pp, root / 'second')
            a, b = Document(first), Document(second)
            self.assertEqual(a.element.xml, b.element.xml)
            self.assertEqual(a.styles.element.xml, b.styles.element.xml)
            audit = self.cli('audit_docx.py', first, second, '--structure-map', mp, '--profile', pp)
            self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)

    def test_table_only_directory_metadata_both_outputs_and_repeat_cli(self):
        from _common import protected_object_manifest
        from test_v051_p3_image_caption_batch_d import ImageCaptionBatchDTests
        with tempfile.TemporaryDirectory(prefix='p3e-directory-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(root)
            profile = json.loads(pp.read_text())
            profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_TABLE_RULE]
            pp.write_text(json.dumps(profile))
            ImageCaptionBatchDTests().add_directory_metadata(source)
            original = source.read_bytes()
            protected = protected_object_manifest(source)
            previous = source
            for name in ('first', 'second'):
                output = self.apply_cli(previous, mp, pp, root / name)
                review = output.with_name(previous.stem + '-review.docx')
                for candidate in (output, review):
                    self.assertEqual(protected, protected_object_manifest(candidate))
                    with zipfile.ZipFile(candidate) as package:
                        for directory in ('word/', 'word/media/', 'word/embeddings/', 'synthetic-empty-directory/'):
                            info = package.getinfo(directory)
                            self.assertTrue(info.is_dir())
                            self.assertEqual(0, info.file_size)
                audit = self.cli('audit_docx.py', previous, output, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)
                if previous != source:
                    self.assertEqual(Document(previous).element.xml, Document(output).element.xml)
                    self.assertEqual(Document(previous).styles.element.xml, Document(output).styles.element.xml)
                previous = output
            self.assertEqual(original, source.read_bytes())

    def test_table_only_missing_real_resource_is_not_restored_cli(self):
        with tempfile.TemporaryDirectory(prefix='p3e-missing-resource-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(root)
            profile = json.loads(pp.read_text())
            profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_TABLE_RULE]
            pp.write_text(json.dumps(profile))
            with zipfile.ZipFile(source, 'a') as package:
                package.writestr('word/media/synthetic-unreferenced.bin', b'nonempty protected resource')
            original = source.read_bytes()
            result = self.cli('apply_profile.py', source, '--structure-map', mp, '--profile', pp,
                              '--output-dir', root / 'out', '--allow-missing-fonts')
            self.assertNotEqual(0, result.returncode)
            self.assertIn('missing_file_not_restored', result.stderr)
            self.assertFalse((root / 'out/synthetic-formatted.docx').exists())
            self.assertEqual(original, source.read_bytes())

    def test_table_only_changed_real_payload_still_fails_product_gate(self):
        import apply_profile
        from test_v051_p3_image_caption_batch_d import PNG, ImageCaptionBatchDTests
        with tempfile.TemporaryDirectory(prefix='p3e-changed-resource-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(
                root, mutate=lambda doc: doc.add_paragraph().add_run().add_picture(io.BytesIO(PNG), width=Inches(1)))
            profile = json.loads(pp.read_text())
            profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_TABLE_RULE]
            pp.write_text(json.dumps(profile))
            helper = ImageCaptionBatchDTests()
            helper.add_directory_metadata(source)
            original = source.read_bytes()
            retain = apply_profile.preserve_zip_directory_metadata
            def corrupt_after_retention(input_path, output_path):
                retain(input_path, output_path)
                with zipfile.ZipFile(output_path) as package:
                    parts = {n: package.read(n) for n in package.namelist()}
                media = next(n for n in parts if n.startswith('word/media/') and not n.endswith('/'))
                parts[media] = b'changed synthetic payload'
                helper.write_package(output_path, parts)
            stdout, stderr = io.StringIO(), io.StringIO()
            argv = ['apply_profile.py', str(source), '--structure-map', str(mp), '--profile', str(pp),
                    '--output-dir', str(root / 'out'), '--allow-missing-fonts']
            with patch.object(sys, 'argv', argv), patch.object(apply_profile, 'preserve_zip_directory_metadata', corrupt_after_retention), contextlib.redirect_stdout(stdout), contextlib.redirect_stderr(stderr):
                code = apply_profile.main()
            self.assertNotEqual(0, code)
            self.assertIn('protected_objects=fail', stderr.getvalue())
            self.assertFalse((root / 'out/synthetic-formatted.docx').exists())
            self.assertEqual(original, source.read_bytes())

    def test_preserves_normal_explicit_row_heights_and_autofit_cli(self):
        from docx.enum.table import WD_ROW_HEIGHT_RULE
        for autofit in (True, False):
            with self.subTest(autofit=autofit), tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
                root = Path(folder)
                def mutate(document):
                    table = document.tables[0]
                    table.autofit = autofit
                    table.rows[0].height, table.rows[0].height_rule = Pt(24), WD_ROW_HEIGHT_RULE.AT_LEAST
                    table.rows[1].height, table.rows[1].height_rule = Pt(24), WD_ROW_HEIGHT_RULE.EXACTLY
                source, mp, pp, _ = self.fixture(root, mutate=mutate)
                output = self.apply_cli(source, mp, pp, root / 'out')
                a, b = Document(source), Document(output)
                self.assertEqual(autofit, b.tables[0].autofit)
                for old, new in zip(a.tables[0].rows, b.tables[0].rows):
                    self.assertEqual(old.height, new.height)
                    self.assertEqual(old.height_rule, new.height_rule)
                self.assertEqual(WD_TABLE_ALIGNMENT.CENTER, b.tables[0].alignment)
                audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)

    def test_full_foundation_image_caption_and_table_rules_coexist_cli(self):
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            root = Path(folder)
            source, mp, pp, _ = self.fixture(root, combined=True)
            output = self.apply_cli(source, mp, pp, root / 'out')
            document = Document(output)
            self.assertEqual(WD_TABLE_ALIGNMENT.CENTER, document.tables[0].alignment)
            self.assertEqual(Pt(6), document.paragraphs[1].paragraph_format.space_before)
            self.assertEqual('Monograph Approved Figure Caption', document.paragraphs[4].style.name)
            self.assertEqual(Pt(18), document.paragraphs[4].paragraph_format.space_after)
            self.assertTrue(document.paragraphs[0].style.name.startswith('Monograph Approved '))
            audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
            self.assertEqual(0, audit.returncode, audit.stdout + audit.stderr)

    def test_audit_rejects_adjacent_note_and_inherited_cell_style_changes(self):
        with tempfile.TemporaryDirectory(prefix='p3e-offline-') as folder:
            root = Path(folder)
            def mutate(document):
                style = document.styles.add_style('Synthetic cell base', 1)
                style.base_style = document.styles['Normal']
                document.tables[0].cell(1, 0).paragraphs[0].style = style
            source, mp, pp, _ = self.fixture(root, mutate=mutate)
            output = self.apply_cli(source, mp, pp, root / 'out')
            saved = output.read_bytes()
            for case in ('note', 'cell_style'):
                with self.subTest(case=case):
                    document = Document(io.BytesIO(saved))
                    if case == 'note':
                        document.paragraphs[2].paragraph_format.space_before = Pt(40)
                    else:
                        document.styles['Synthetic cell base'].font.bold = True
                    document.save(output)
                    audit = self.cli('audit_docx.py', source, output, '--structure-map', mp, '--profile', pp)
                    self.assertNotEqual(0, audit.returncode, audit.stdout)


if __name__ == '__main__':
    unittest.main()
