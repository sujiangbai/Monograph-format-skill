from __future__ import annotations

import base64
import copy
import io
import json
import sys
import subprocess
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'format-monograph/scripts'))
import structure_map as sm
from _common import apply_rule
from _common import FormatMonographError

PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
RULE = {'id': 'FMT-FIGCAP-501', 'selector': {'kind': 'caption_role', 'value': 'figure_caption'},
        'properties': {'font_size_pt': 9, 'alignment': 'center', 'line_spacing_rule': 'at_least',
                       'line_spacing_pt': 15, 'space_before_pt': 0}}


class ImageCaptionBatchDTests(unittest.TestCase):
    def fixture(self, following_table=False, cropped=False, old_roles=False, unnumbered=False):
        document = Document()
        document.sections[0].header.paragraphs[0].text = 'Existing header'
        document.sections[0].footer.paragraphs[0].text = 'Existing footer'
        image = document.add_paragraph()
        image.add_run().add_picture(io.BytesIO(PNG), width=Inches(1))
        if cropped:
            crop = OxmlElement('a:srcRect')
            crop.set('l', '5000')
            image._p.xpath('.//pic:blipFill')[0].append(crop)
        image.paragraph_format.left_indent = Pt(20)
        caption = document.add_paragraph('Existing unnumbered caption' if unnumbered else '图 1-1 Synthetic caption', 'Caption')
        document.add_paragraph('Synthetic following body')
        document.add_paragraph('Unapproved table caption', 'Caption')
        if old_roles:
            document.add_paragraph('Existing book title', 'Title')
            document.add_paragraph('Existing chapter', 'Heading 1')
        if following_table:
            table = document.add_table(rows=1, cols=1)
            table.cell(0, 0).text = 'Existing data table'
            caption._p.addnext(table._tbl)
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / 'synthetic.docx'
            document.save(source)
            mapping = sm.candidate_structure_map(source)
            self.source_bytes = source.read_bytes()
        mapping['status'] = 'approved'
        if unnumbered:
            role = mapping['paragraph_roles'][0]
            role.update(role='figure_caption_unnumbered', canonical_role='figure_caption_unnumbered')
            mapping['pagination_groups'] = [{'kind': 'figure_with_caption',
                'anchor': copy.deepcopy(mapping['images'][0]['locator']),
                'caption': copy.deepcopy(role['locator']), 'approved': False}]
        if following_table:
            mapping['tables'][0].update(approved=True, kind='data')
        for entry in mapping['paragraph_roles']:
            if entry['locator']['paragraph'] in {1, 2}:
                entry['approved'] = True
            if entry['locator']['paragraph'] == 2:
                entry.update(role='body_text', canonical_role='body_text')
            if old_roles and entry['locator']['paragraph'] in {4, 5}:
                entry['approved'] = True
        mapping['pagination_groups'][0]['approved'] = True
        for entry in mapping['captions']:
            entry.update(approved=True, action='style_only')
        return document, mapping

    def apply(self, document, mapping):
        return sm._apply_foundation_figure_captions(document, mapping, RULE)

    def test_independent_image_paragraph_format(self):
        document, mapping = self.fixture()
        self.apply(document, mapping)
        self.assertEqual(Pt(6), document.paragraphs[0].paragraph_format.space_before)
        self.assertEqual(Pt(0), document.paragraphs[0].paragraph_format.left_indent)

    def test_unapproved_shared_caption_style_unchanged(self):
        document, mapping = self.fixture()
        before = document.styles['Caption'].element.xml
        self.apply(document, mapping)
        self.assertEqual(before, document.styles['Caption'].element.xml)

    def test_context_matrix(self):
        for role, expected in [('body_text', 18), ('list_item', 18), ('chapter_title', 0),
                               ('image', 0), ('display_equation', 0), ('page_break', 0)]:
            with self.subTest(role=role):
                document, mapping = self.fixture()
                following = document.paragraphs[2]
                entry = mapping['paragraph_roles'][1]
                entry.update(role=role, canonical_role=role)
                if role == 'image':
                    following.add_run().add_picture(io.BytesIO(PNG), width=Inches(1))
                if role == 'display_equation':
                    following._p.append(OxmlElement('m:oMathPara'))
                if role == 'page_break':
                    following.paragraph_format.page_break_before = True
                result = self.apply(document, mapping)
                self.assertEqual(1, result['pairs'])
                self.assertEqual(Pt(expected), document.paragraphs[1].paragraph_format.space_after)
                self.assertEqual('pending_visual_verification', result['same_page'])

    def test_unapproved_and_out_of_scope_preserved(self):
        for case in ['pair', 'image', 'caption', 'following', 'mixed', 'multiple', 'floating', 'note']:
            with self.subTest(case=case):
                document, mapping = self.fixture()
                if case == 'pair': mapping['pagination_groups'][0]['approved'] = False
                if case == 'image': mapping['images'] = []
                if case == 'caption': mapping['paragraph_roles'][0]['approved'] = False
                if case == 'following': mapping['paragraph_roles'][1]['approved'] = False
                if case == 'mixed': document.paragraphs[0].add_run('authored text')
                if case == 'multiple': document.paragraphs[0].add_run().add_picture(io.BytesIO(PNG))
                if case == 'floating': document.paragraphs[0]._p.xpath('.//wp:inline')[0].tag = qn('wp:anchor')
                if case == 'note': mapping['paragraph_roles'][1].update(role='figure_note', canonical_role='figure_note')
                before = document.element.xml
                result = self.apply(document, mapping)
                self.assertEqual(0, result['pairs'])
                self.assertTrue(result['skipped'])
                self.assertEqual(before, document.element.xml)

    def test_following_table_preserved(self):
        document, mapping = self.fixture()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Existing data'
        document.paragraphs[1]._p.addnext(table._tbl)
        mapping['tables'] = [{'table': 0, 'approved': True, 'kind': 'data',
                              'table_text_sha256': sm._table_text_hash(table)}]
        before = table._tbl.xml
        result = self.apply(document, mapping)
        self.assertEqual(1, result['pairs'])
        self.assertEqual(Pt(0), document.paragraphs[1].paragraph_format.space_after)
        self.assertEqual(before, table._tbl.xml)

    def test_in_cell_pair_and_duplicate_role_are_reported(self):
        for case in ['cell', 'duplicate_role']:
            with self.subTest(case=case):
                document, mapping = self.fixture()
                if case == 'cell':
                    mapping['pagination_groups'][0]['anchor']['kind'] = 'table_cell_paragraph'
                else:
                    mapping['paragraph_roles'].append(copy.deepcopy(mapping['paragraph_roles'][0]))
                before = document.element.xml
                self.assertEqual(0, self.apply(document, mapping)['pairs'])
                self.assertEqual(before, document.element.xml)

    def test_readonly_data_and_layout_successors_cli(self):
        for kind in ('data', 'layout'):
            with self.subTest(kind=kind):
                self._real_apply_audit_cli(following_table=True, readonly_table=kind)

    def test_readonly_table_identity_negative_matrix(self):
        for case in ('pair_unapproved', 'wrong_index', 'duplicate', 'wrong_hash', 'not_adjacent', 'string_index'):
            with self.subTest(case=case):
                document, mapping = self.fixture(following_table=True)
                mapping['tables'][0].update(approved=False, kind='layout')
                if case == 'pair_unapproved': mapping['pagination_groups'][0]['approved'] = False
                if case == 'wrong_index': mapping['tables'][0]['table'] = 99
                if case == 'duplicate': mapping['tables'].append(copy.deepcopy(mapping['tables'][0]))
                if case == 'wrong_hash': mapping['tables'][0]['table_text_sha256'] = '0' * 64
                if case == 'string_index': mapping['tables'][0]['table'] = '0'
                if case == 'not_adjacent':
                    spacer = document.add_paragraph()
                    document.paragraphs[1]._p.addnext(spacer._p)
                before = document.element.xml
                if case == 'wrong_hash':
                    with self.assertRaises(FormatMonographError): self.apply(document, mapping)
                else:
                    self.assertEqual(0, self.apply(document, mapping)['pairs'])
                self.assertEqual(before, document.element.xml)

    def test_approved_layout_operation_gate_not_weakened(self):
        _, mapping = self.fixture(following_table=True)
        mapping['tables'][0].update(kind='layout', approved=True, pagination_only=False)
        mapping['tables'][0]['visual'] = {'approved': False}
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'map.json'
            path.write_text(json.dumps(mapping))
            with self.assertRaisesRegex(FormatMonographError, 'Approved layout tables'):
                sm.load_structure_map(path)
            mapping['tables'][0]['approved'] = False
            path.write_text(json.dumps(mapping))
            sm.load_structure_map(path)

    def test_legacy_default_qa_and_formula_image_block(self):
        from apply_profile import preflight_equations
        from unittest.mock import patch
        profile = json.loads((ROOT / 'format-monograph/examples/profiles/v051-foundation-format-slice.json').read_text())
        self.assertEqual('qa', profile['runtime_policy']['legacy_equation_policy'])
        inventory = {'formula_image_candidates': 0, 'legacy_equation_ole': 1}
        with patch('apply_profile.equation_inventory', return_value=inventory):
            with self.assertRaisesRegex(FormatMonographError, 'Legacy Equation Editor'):
                preflight_equations(Path('synthetic.docx'), profile)
            profile['runtime_policy']['legacy_equation_policy'] = 'allow'
            self.assertEqual(inventory, preflight_equations(Path('synthetic.docx'), profile))
            inventory['formula_image_candidates'] = 1
            with self.assertRaisesRegex(FormatMonographError, 'Editable-equation policy'):
                preflight_equations(Path('synthetic.docx'), profile)
        self.assertTrue(profile['runtime_policy']['editable_equations_required'])
        self.assertEqual('block', profile['runtime_policy']['formula_image_policy'])

    def test_legacy_allow_cli_preserves_object_and_rejects_mutation(self):
        self._real_apply_audit_cli(legacy=True)

    def test_real_formula_image_still_blocked_under_allow(self):
        from apply_profile import preflight_equations
        document, _ = self.fixture()
        document.paragraphs[0]._p.xpath('.//wp:docPr')[0].set('descr', 'Synthetic formula image')
        profile = json.loads((ROOT / 'format-monograph/examples/profiles/v051-foundation-format-slice.json').read_text())
        profile['runtime_policy']['legacy_equation_policy'] = 'allow'
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'synthetic.docx'
            document.save(path)
            with self.assertRaisesRegex(FormatMonographError, 'Editable-equation policy'):
                preflight_equations(path, profile)

    def test_unauthorized_mutation_rejected(self):
        for case in ['caption_text', 'extent', 'reorder', 'resize', 'duplicate_pair', 'renumber', 'blank_insertion']:
            with self.subTest(case=case):
                document, mapping = self.fixture()
                if case == 'caption_text': document.paragraphs[1].add_run(' changed')
                if case == 'extent': document.paragraphs[0]._p.xpath('.//wp:extent')[0].set('cx', '123')
                if case == 'reorder': document.paragraphs[1]._p.addprevious(document.paragraphs[2]._p)
                if case == 'resize': mapping['images'][0]['resize']['approved'] = True
                if case == 'duplicate_pair': mapping['pagination_groups'] *= 2
                if case == 'renumber': mapping['captions'][0]['action'] = 'replace_identifier'
                if case == 'blank_insertion': mapping['block_spacing']['approved'] = True
                with self.assertRaises(FormatMonographError): self.apply(document, mapping)

    def test_direct_spacing_overrides_are_removed(self):
        document, mapping = self.fixture()
        for paragraph in document.paragraphs[:2]:
            spacing = paragraph._p.get_or_add_pPr().get_or_add_spacing()
            for attribute in ['beforeLines', 'afterLines', 'beforeAutospacing', 'afterAutospacing']:
                spacing.set(qn('w:' + attribute), '100')
        self.apply(document, mapping)
        for paragraph in document.paragraphs[:2]:
            spacing = paragraph._p.pPr.find(qn('w:spacing'))
            for attribute in ['beforeLines', 'afterLines', 'beforeAutospacing', 'afterAutospacing']:
                self.assertIsNone(spacing.get(qn('w:' + attribute)))

    def test_protected_payload_and_semantic_idempotence(self):
        document, mapping = self.fixture()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = 'Unapproved table'
        math = OxmlElement('m:oMathPara')
        document.add_paragraph()._p.append(math)
        before_table, before_math = table._tbl.xml, etree.tostring(math)
        drawing = document.paragraphs[0]._p.xpath('.//w:drawing')[0]
        before_drawing = drawing.xml
        before_text = [p.text for p in document.paragraphs]
        before_order = list(document.element.body)
        self.apply(document, mapping)
        first = document.element.xml
        self.apply(document, mapping)
        self.assertEqual(first, document.element.xml)
        self.assertEqual(before_drawing, drawing.xml)
        self.assertEqual(before_table, table._tbl.xml)
        self.assertEqual(before_math, etree.tostring(math))
        self.assertEqual(before_text, [p.text for p in document.paragraphs])
        self.assertEqual(before_order, list(document.element.body))

    def test_real_apply_audit_cli(self):
        self._real_apply_audit_cli()

    def correction_inputs(self, root, *, style_case=None, outline=False, heading=False, combined=False, duplicate_media=False, skipped=False, adjacent=False):
        from docx.opc.part import Part
        from docx.opc.packuri import PackURI
        self.fixture()
        doc=Document(io.BytesIO(self.source_bytes))
        if outline:
            node=OxmlElement('w:outlineLvl');node.set(qn('w:val'),'0')
            doc.paragraphs[2]._p.get_or_add_pPr().append(node)
        if heading:doc.paragraphs[2].style=doc.styles['Heading 1']
        if skipped:doc.paragraphs[0]._p.xpath('.//wp:inline')[0].tag=qn('wp:anchor')
        if adjacent:
            doc.paragraphs[2].clear().add_run().add_picture(io.BytesIO(PNG),width=Inches(1))
            doc.paragraphs[3].text='图 1-2 Second synthetic caption'
            doc.add_paragraph('Second pair following body')
        if style_case:
            name='Monograph Approved Figure Caption';identifier=name.replace(' ','')
            if not style_case.startswith('dangling'):
                derived=doc.styles.add_style(name,1);derived.base_style=doc.styles['Caption']
                doc.paragraphs[1].style=derived
            if style_case in ('header','dangling_header','body','dangling_body'):
                p=doc.sections[0].header.paragraphs[0] if 'header' in style_case else doc.paragraphs[3]
                p._p.get_or_add_pPr().get_or_add_pStyle().val=identifier
            if style_case in ('inheritance','dangling_inheritance'):
                child=doc.styles.add_style('Synthetic descendant',1)
                etree.SubElement(child.element,qn('w:basedOn')).set(qn('w:val'),identifier)
        if duplicate_media:
            drawing=doc.paragraphs[0]._p.xpath('.//w:drawing')[0]
            rid=drawing.xpath('.//a:blip/@r:embed')[0];old=doc.part.rels[rid].target_part
            extra=Part(PackURI('/word/media/synthetic-duplicate.png'),old.content_type,old.blob,doc.part.package)
            doc.part.relate_to(extra,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/image')
        source=root/'synthetic.docx';doc.save(source)
        mapping=sm.candidate_structure_map(source);mapping['status']='approved'
        for g in mapping['pagination_groups']:
            if g['kind']=='figure_with_caption' and g['caption'].get('paragraph') in ({1,3} if adjacent else {1}):g['approved']=True
        for e in mapping['paragraph_roles']:
            if e['locator'].get('paragraph')==1:e.update(approved=True,role='figure_caption',canonical_role='figure_caption')
            if e['locator'].get('paragraph')==2:e.update(approved=True,role='body_text',canonical_role='body_text')
            if adjacent and e['locator'].get('paragraph')==2:e.update(approved=True,role='image',canonical_role='image')
            if adjacent and e['locator'].get('paragraph')==3:e.update(approved=True,role='figure_caption',canonical_role='figure_caption')
            if adjacent and e['locator'].get('paragraph')==4:e.update(approved=True,role='body_text',canonical_role='body_text')
        for e in mapping['captions']:
            if e['locator'].get('paragraph') in ({1,3} if adjacent else {1}):e.update(approved=True,action='style_only')
        profile=json.loads((ROOT/'format-monograph/examples/profiles/v051-foundation-format-slice.json').read_text())
        profile['rules']=[r for r in profile['rules'] if r['id']==sm.FOUNDATION_FIGURE_RULE or (combined and r['selector']=={'kind':'paragraph_role','value':'body_text'})]
        self.assertEqual(2 if combined else 1,len(profile['rules']))
        mp=root/'map.json';mp.write_text(json.dumps(mapping));pp=root/'profile.json';pp.write_text(json.dumps(profile))
        return source,mp,pp

    def correction_cli(self, script, *args):
        return subprocess.run([sys.executable,'-B',str(ROOT/'format-monograph/scripts'/script),*map(str,args)],capture_output=True,text=True)

    def test_caption_non_target_story_dangling_and_inheritance_refused_before_write(self):
        for case in ('header','dangling_header','body','dangling_body','inheritance','dangling_inheritance'):
            with self.subTest(case=case),tempfile.TemporaryDirectory() as folder:
                root=Path(folder);source,mp,pp=self.correction_inputs(root,style_case=case)
                before=source.read_bytes()
                result=self.correction_cli('apply_profile.py',source,'--structure-map',mp,'--profile',pp,'--output-dir',root/'out','--allow-missing-fonts')
                self.assertNotEqual(0,result.returncode,result.stdout)
                self.assertIn('non-target or inheritance reference',result.stderr)
                self.assertEqual(before,source.read_bytes())
                self.assertFalse(list((root/'out').glob('*.docx')))

    def test_actual_image_relationship_identity_audit_matrix(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source,mp,pp=self.correction_inputs(root,duplicate_media=True)
            applied=self.correction_cli('apply_profile.py',source,'--structure-map',mp,'--profile',pp,'--output-dir',root/'out','--allow-missing-fonts')
            self.assertEqual(0,applied.returncode,applied.stderr)
            formatted=root/'out/synthetic-formatted.docx'
            with zipfile.ZipFile(formatted) as z:parts={n:z.read(n) for n in z.namelist()}
            doc=etree.fromstring(parts['word/document.xml']);rid=doc.xpath('.//a:blip/@r:embed',namespaces={
                'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
                'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'})[0]
            for case in ('same_target_equivalent_path','same_bytes_different_member','type_changed','external_mode'):
                modified=dict(parts);rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
                rel=next(n for n in rels if n.get('Id')==rid)
                if case=='same_target_equivalent_path':rel.set('Target','./media/../'+rel.get('Target'))
                if case=='same_bytes_different_member':rel.set('Target','media/synthetic-duplicate.png')
                if case=='type_changed':rel.set('Type','http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink')
                if case=='external_mode':rel.set('TargetMode','External')
                modified['word/_rels/document.xml.rels']=etree.tostring(rels);self.write_package(formatted,modified)
                audited=self.correction_cli('audit_docx.py',source,formatted,'--structure-map',mp,'--profile',pp)
                self.assertEqual(case=='same_target_equivalent_path',audited.returncode==0,case+audited.stderr+audited.stdout)

    def test_context_outline_readonly_and_independent_body_rule_cli(self):
        for heading,combined in ((False,False),(True,False),(True,True)):
            with self.subTest(heading=heading,combined=combined),tempfile.TemporaryDirectory() as folder:
                root=Path(folder);source,mp,pp=self.correction_inputs(root,outline=True,heading=heading,combined=combined)
                applied=self.correction_cli('apply_profile.py',source,'--structure-map',mp,'--profile',pp,'--output-dir',root/'out','--allow-missing-fonts')
                self.assertEqual(0,applied.returncode,applied.stderr)
                formatted=root/'out/synthetic-formatted.docx';a,b=Document(source),Document(formatted)
                if combined:
                    self.assertFalse(b.paragraphs[2]._p.xpath('./w:pPr/w:outlineLvl'))
                    self.assertTrue(b.paragraphs[2].style.name.startswith('Monograph Approved'))
                else:
                    self.assertEqual(a.paragraphs[2]._p.xml,b.paragraphs[2]._p.xml)
                audited=self.correction_cli('audit_docx.py',source,formatted,'--structure-map',mp,'--profile',pp)
                self.assertEqual(0,audited.returncode,audited.stderr+audited.stdout)
                if not combined:
                    b.paragraphs[2]._p.pPr.remove(b.paragraphs[2]._p.pPr.find(qn('w:outlineLvl')));b.save(formatted)
                    rejected=self.correction_cli('audit_docx.py',source,formatted,'--structure-map',mp,'--profile',pp)
                    self.assertNotEqual(0,rejected.returncode)

    def test_adjacent_pairs_and_skipped_pair_context_cli(self):
        for adjacent in (True,False):
            with self.subTest(adjacent=adjacent),tempfile.TemporaryDirectory() as folder:
                root=Path(folder);source,mp,pp=self.correction_inputs(root,adjacent=adjacent,skipped=not adjacent,outline=not adjacent)
                mapping=sm.load_structure_map(mp)
                self.assertEqual(2 if adjacent else 1,sum(g.get('approved',False) for g in mapping['pagination_groups']))
                applied=self.correction_cli('apply_profile.py',source,'--structure-map',mp,'--profile',pp,'--output-dir',root/'out','--allow-missing-fonts')
                self.assertEqual(0,applied.returncode,applied.stderr)
                formatted=root/'out/synthetic-formatted.docx';a,b=Document(source),Document(formatted)
                if adjacent:
                    self.assertNotEqual(a.paragraphs[2]._p.xml,b.paragraphs[2]._p.xml)
                    self.assertEqual('Monograph Approved Figure Caption',b.paragraphs[3].style.name)
                else:
                    self.assertEqual(a.paragraphs[2]._p.xml,b.paragraphs[2]._p.xml)
                audited=self.correction_cli('audit_docx.py',source,formatted,'--structure-map',mp,'--profile',pp)
                self.assertEqual(0,audited.returncode,audited.stderr+audited.stdout)

    def test_missing_caption_locator_skipped_apply_and_audit_cli(self):
        with tempfile.TemporaryDirectory() as folder:
            root=Path(folder);source,mp,pp=self.correction_inputs(root)
            mapping=json.loads(mp.read_text())
            group=next(g for g in mapping['pagination_groups'] if g.get('approved'))
            del group['caption'];mp.write_text(json.dumps(mapping))
            checked=sm.load_structure_map(mp)
            rule=json.loads(pp.read_text())['rules'][0]
            plans,skipped=sm._foundation_figure_plan(Document(source),checked,rule)
            self.assertEqual([],plans)
            self.assertIn('not_body_pair',[e.get('reason') for e in skipped])
            self.assertEqual(set(),sm._foundation_readonly_contexts(Document(source),checked,plans,[rule]))
            original=source.read_bytes()
            applied=self.correction_cli('apply_profile.py',source,'--structure-map',mp,'--profile',pp,'--output-dir',root/'out','--allow-missing-fonts')
            self.assertEqual(0,applied.returncode,applied.stderr)
            self.assertNotIn('Traceback',applied.stderr+applied.stdout)
            report=(root/'out/synthetic-format-report.md').read_text()
            evidence=next(json.loads(line) for line in report.splitlines() if line.startswith('{') and '"pairs"' in line)
            self.assertEqual(0,evidence['pairs'])
            self.assertIn('not_body_pair',[e.get('reason') for e in evidence['skipped']])
            formatted=root/'out/synthetic-formatted.docx'
            audited=self.correction_cli('audit_docx.py',source,formatted,'--structure-map',mp,'--profile',pp)
            self.assertEqual(0,audited.returncode,audited.stderr+audited.stdout)
            self.assertNotIn('Traceback',audited.stderr+audited.stdout)
            self.assertEqual(original,source.read_bytes())
            self.assertEqual([],sm._foundation_figure_plan(Document(formatted),checked,rule)[0])

    def add_directory_metadata(self, source):
        with zipfile.ZipFile(source) as z:
            parts = {n: z.read(n) for n in z.namelist()}
        parts['word/'] = b''
        parts['word/media/'] = b''
        parts['word/embeddings/'] = b''
        parts['synthetic-empty-directory/'] = b''
        self.write_package(source, parts)
        return parts

    def write_package(self, path, parts):
        with zipfile.ZipFile(path, 'w') as z:
            for name, data in parts.items(): z.writestr(name, data)

    def test_directory_retention_both_outputs_and_repeat_cli(self):
        self._real_apply_audit_cli(unreferenced=True)

    def test_directory_preservation_negative_matrix(self):
        from apply_profile import preserve_zip_directory_metadata
        from _common import protected_object_manifest
        for case in ('missing_protected_file', 'changed_protected_file', 'other_file', 'nonempty_directory'):
            with self.subTest(case=case), tempfile.TemporaryDirectory() as directory:
                source, output = [Path(directory)/n for n in ('source.docx', 'output.docx')]
                self.fixture(); source.write_bytes(self.source_bytes)
                parts = self.add_directory_metadata(source)
                Document(source).save(output)
                with zipfile.ZipFile(output) as z: saved={n:z.read(n) for n in z.namelist()}
                target=next(n for n in parts if n.startswith('word/media/') and not n.endswith('/'))
                if case=='missing_protected_file': del saved[target]
                if case=='changed_protected_file': saved[target]=b'changed protected bytes'
                if case=='other_file': parts['custom/untouched.bin']=b'out of scope'
                if case=='nonempty_directory': parts['synthetic-empty-directory/']=b'not metadata'
                self.write_package(source,parts);self.write_package(output,saved)
                if case=='changed_protected_file':
                    preserve_zip_directory_metadata(source,output)
                    self.assertNotEqual(protected_object_manifest(source),protected_object_manifest(output))
                    with zipfile.ZipFile(output) as z:self.assertEqual(saved[target],z.read(target))
                else:
                    reason='nonempty_directory' if case=='nonempty_directory' else 'missing_file_not_restored'
                    with self.assertRaisesRegex(FormatMonographError,'reason='+reason):preserve_zip_directory_metadata(source,output)
                    self.assertFalse(output.exists())
                self.assertFalse(list(Path(directory).glob('.p3d-retain-*')))


    def test_real_apply_audit_cli_following_table(self):
        self._real_apply_audit_cli(following_table=True)

    def test_unnumbered_caption_cli_does_not_generate_number(self):
        self._real_apply_audit_cli(unnumbered=True)

    def test_complete_foundation_cli_with_existing_crop_and_old_roles(self):
        self._real_apply_audit_cli(cropped=True, full_profile=True)

    def _real_apply_audit_cli(self, following_table=False, cropped=False, full_profile=False, unnumbered=False, readonly_table=None, legacy=False, unreferenced=False):
        document, mapping = self.fixture(following_table=following_table, cropped=cropped, old_roles=full_profile, unnumbered=unnumbered)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            source = root / 'synthetic.docx'
            source.write_bytes(self.source_bytes)
            if unreferenced:
                from _common import protected_object_manifest
                self.add_directory_metadata(source)
                control=root/'control.docx'
                Document(source).save(control)
                self.assertNotEqual(protected_object_manifest(source), protected_object_manifest(control))
            if readonly_table:
                mapping['tables'][0].update(approved=False, kind=readonly_table)
            if legacy:
                from docx.opc.part import Part
                from docx.opc.packuri import PackURI
                doc = Document(source)
                part = Part(PackURI('/word/embeddings/synthetic.bin'), 'application/vnd.openxmlformats-officedocument.oleObject', b'SYNTHETIC INERT OBJECT', doc.part.package)
                rel = doc.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject')
                obj = OxmlElement('w:object')
                ole = etree.SubElement(obj, '{urn:schemas-microsoft-com:office:office}OLEObject', ProgID='Equation.3')
                ole.set(qn('r:id'), rel)
                doc.add_paragraph().add_run()._r.append(obj)
                doc.save(source)
                mapping['source_content_fingerprint_sha256'] = sm.content_fingerprint(source)
            map_path = root / 'map.json'
            map_path.write_text(json.dumps(mapping))
            profile = json.loads((ROOT / 'format-monograph/examples/profiles/v051-foundation-format-slice.json').read_text())
            if not full_profile:
                profile['rules'] = [r for r in profile['rules'] if r['id'] == sm.FOUNDATION_FIGURE_RULE]
            if legacy:
                from apply_profile import preflight_equations
                with self.assertRaisesRegex(FormatMonographError, 'Legacy Equation Editor'):
                    preflight_equations(source, profile)
                profile['runtime_policy']['legacy_equation_policy'] = 'allow'
            profile_path = root / 'profile.json'
            profile_path.write_text(json.dumps(profile))
            def run(script, *args):
                return subprocess.run([sys.executable, '-B', str(ROOT / 'format-monograph/scripts' / script), *map(str, args)], capture_output=True, text=True)
            applied = run('apply_profile.py', source, '--profile', profile_path, '--structure-map', map_path,
                          '--output-dir', root / 'out', '--allow-missing-fonts')
            self.assertEqual(0, applied.returncode, applied.stdout + applied.stderr)
            formatted = root / 'out/synthetic-formatted.docx'
            if unreferenced:
                repeated = run('apply_profile.py', formatted, '--profile', profile_path, '--structure-map', map_path,
                               '--output-dir', root/'repeat', '--allow-missing-fonts')
                self.assertEqual(0, repeated.returncode, repeated.stdout+repeated.stderr)
                for candidate in (formatted, root/'out/synthetic-review.docx',
                                  root/'repeat/synthetic-formatted-formatted.docx', root/'repeat/synthetic-formatted-review.docx'):
                    self.assertEqual(protected_object_manifest(source), protected_object_manifest(candidate))
                    with zipfile.ZipFile(candidate) as z:
                        for name in ('word/', 'word/media/', 'word/embeddings/', 'synthetic-empty-directory/'):
                            info=z.getinfo(name)
                            self.assertTrue(info.is_dir())
                            self.assertEqual(0,info.file_size)
                        self.assertEqual(len(z.namelist()), len(set(z.namelist())))
                        types=etree.fromstring(z.read('[Content_Types].xml'))
                        keys=[(n.tag,n.get('PartName',n.get('Extension'))) for n in types]
                        self.assertEqual(len(keys),len(set(keys)))
                with zipfile.ZipFile(formatted) as a, zipfile.ZipFile(root/'repeat/synthetic-formatted-formatted.docx') as b:
                    self.assertEqual(set(a.namelist()),set(b.namelist()))
                    for n in a.namelist():
                        if n.endswith(('.xml','.rels')):
                            self.assertEqual(etree.tostring(etree.fromstring(a.read(n)),method='c14n'),
                                             etree.tostring(etree.fromstring(b.read(n)),method='c14n'))
                        else: self.assertEqual(a.read(n),b.read(n))
            with zipfile.ZipFile(source) as before, zipfile.ZipFile(formatted) as after:
                for name in before.namelist():
                    if name.startswith(('word/media/', 'word/embeddings/', 'word/header', 'word/footer', 'word/_rels/')):
                        self.assertEqual(before.read(name), after.read(name), name)
            original_doc, formatted_doc = Document(source), Document(formatted)
            self.assertEqual([p.text for p in original_doc.paragraphs], [p.text for p in formatted_doc.paragraphs])
            self.assertEqual(original_doc.paragraphs[0]._p.xpath('.//w:drawing')[0].xml,
                             formatted_doc.paragraphs[0]._p.xpath('.//w:drawing')[0].xml)
            if full_profile:
                self.assertTrue(formatted_doc.paragraphs[4].style.name.startswith('Monograph Approved'))
                self.assertTrue(formatted_doc.paragraphs[5].style.name.startswith('Monograph Approved'))
                self.assertEqual('Caption', formatted_doc.paragraphs[3].style.name)
            self.assertEqual(original_doc.sections[0]._sectPr.xml, formatted_doc.sections[0]._sectPr.xml)
            if following_table:
                self.assertEqual(original_doc.tables[0]._tbl.xml, formatted_doc.tables[0]._tbl.xml)
                self.assertEqual(list(original_doc.element.body).index(original_doc.tables[0]._tbl),
                                 list(formatted_doc.element.body).index(formatted_doc.tables[0]._tbl))
            audited = run('audit_docx.py', source, formatted, '--profile', profile_path, '--structure-map', map_path)
            self.assertEqual(0, audited.returncode, audited.stdout + audited.stderr)
            report = (root / 'out/synthetic-format-report.md').read_text()
            self.assertIn('pending_visual_verification', report)
            good = formatted.read_bytes()
            if legacy:
                self.assertEqual([etree.tostring(x, method='c14n') for x in original_doc.element.xpath('.//w:object')],
                                 [etree.tostring(x, method='c14n') for x in formatted_doc.element.xpath('.//w:object')])
                with zipfile.ZipFile(io.BytesIO(good)) as package:
                    parts = {name: package.read(name) for name in package.namelist()}
                parts['word/embeddings/synthetic.bin'] = b'UNAUTHORIZED OBJECT CHANGE'
                with zipfile.ZipFile(formatted, 'w') as package:
                    for name, value in parts.items(): package.writestr(name, value)
                audited = run('audit_docx.py', source, formatted, '--profile', profile_path, '--structure-map', map_path)
                self.assertNotEqual(0, audited.returncode)
            for case in ['spacing', 'text', 'extent', 'image_becomes_mixed'] + (['crop'] if cropped else []):
                with self.subTest(tamper=case):
                    changed = Document(io.BytesIO(good))
                    if case == 'spacing': changed.paragraphs[1].paragraph_format.space_after = Pt(18 if following_table else 0)
                    if case == 'text': changed.paragraphs[1].add_run(' unauthorized')
                    if case == 'extent': changed.paragraphs[0]._p.xpath('.//wp:extent')[0].set('cx', '300')
                    if case == 'image_becomes_mixed': changed.paragraphs[0].add_run(' unauthorized')
                    if case == 'crop': changed.paragraphs[0]._p.xpath('.//a:srcRect')[0].set('l', '10000')
                    changed.save(formatted)
                    audited = run('audit_docx.py', source, formatted, '--profile', profile_path, '--structure-map', map_path)
                    self.assertNotEqual(0, audited.returncode, audited.stdout + audited.stderr)


if __name__ == '__main__':
    unittest.main()
